"""Per-catalog / per-merge_method dataset counts from dublin_core_metadata.

Old side  : one old dataset == one row (nid is unique); `old_catalogs` counts the
            distinct Relation[Catalog Title] values -- the pre-merge catalogs --
            folded into the new catalog.
New side  : mergeable rows collapse into their Collection; rows without a
            Collection (non-mergeable) stay standalone datasets; rows flagged
            'to be removed' are dropped, so they contribute 0 new datasets.

Collection -- not new_title -- is authoritative for the merged name: the
`curate datasets` xlsx update rewrote Collection on 30,802 curate rows and
left new_title stale.

Writes the same report as CSV and as a two-sheet xlsx; the second sheet is the
old-catalog -> new-catalog mapping, one row per old catalog.
"""

import csv
import os

import duckdb
import openpyxl

DB = os.path.join(os.path.dirname(__file__), "..", "metadata.db")
DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data")
OUT = os.path.join(DATA_DIR, "catalog_merge_method_stats.csv")
OUT_XLSX = os.path.join(DATA_DIR, "catalog_merge_method_stats.xlsx")

# Rows whose merge_method means the dataset is retired rather than republished.
DROPPED_METHODS = ("to be removed",)

STATS_SQL = """
WITH base AS (
    SELECT
        new_catalog_title,
        merge_method,
        nid,
        "Relation[Catalog Title]" AS old_catalog,
        NULLIF(trim(Collection), '') AS coll,
        merge_method IN ('to be removed') AS is_dropped
    FROM dublin_core_metadata
)
SELECT
    new_catalog_title,
    merge_method,
    count(DISTINCT old_catalog)                          AS old_catalogs,
    count(*)                                             AS old_datasets,
    CASE WHEN bool_and(is_dropped) THEN 0 ELSE
        count(DISTINCT coll)
        + sum(CASE WHEN coll IS NULL AND NOT is_dropped THEN 1 ELSE 0 END)
    END                                                  AS new_datasets
FROM base
GROUP BY 1, 2
ORDER BY 1, old_datasets DESC
"""

# An old catalog can appear under several merge_methods, so catalog TOTAL and
# GRAND TOTAL rows need their own DISTINCT counts -- summing the method rows
# double-counts (HMIS: 161 summed vs 151 distinct).
OLD_CATALOGS_SQL = """
SELECT new_catalog_title, count(DISTINCT "Relation[Catalog Title]")
FROM dublin_core_metadata GROUP BY 1
"""

OLD_CATALOGS_ALL_SQL = """
SELECT count(DISTINCT "Relation[Catalog Title]") FROM dublin_core_metadata
"""

# One row per old catalog. No old catalog spans two new catalogs (verified: the
# per-catalog distincts sum to the global 365), so new_catalog_title is single
# valued here; merge_method is not -- 35 old catalogs carry several -- so those
# collapse into one comma-separated cell, busiest method first.
MAPPING_SQL = """
WITH per_method AS (
    SELECT
        "Relation[Catalog Title]" AS old_catalog,
        new_catalog_title,
        merge_method,
        count(*) AS n
    FROM dublin_core_metadata
    GROUP BY 1, 2, 3
)
SELECT
    old_catalog,
    new_catalog_title,
    string_agg(merge_method, ', ' ORDER BY n DESC, merge_method) AS merge_methods,
    sum(n) AS old_datasets
FROM per_method
GROUP BY 1, 2
ORDER BY old_datasets DESC, old_catalog
"""


def main():
    con = duckdb.connect(DB, read_only=True)
    rows = con.execute(STATS_SQL).fetchall()
    old_catalogs_total = dict(con.execute(OLD_CATALOGS_SQL).fetchall())
    old_catalogs_grand = con.execute(OLD_CATALOGS_ALL_SQL).fetchone()[0]

    # Collection never spans two merge_methods (verified), so a catalog TOTAL is
    # just the sum of its method rows -- except old_catalogs, which does span.
    by_catalog = {}
    for cat, *_ in rows:
        by_catalog.setdefault(cat, [])
    for r in rows:
        by_catalog[r[0]].append(r)

    header = ["new_catalog_title", "merge_method", "old_catalogs",
              "old_datasets", "new_datasets"]

    out_rows = []
    for cat in sorted(by_catalog, key=lambda c: -sum(r[3] for r in by_catalog[c])):
        group = by_catalog[cat]
        for _, mm, oldc, old, new in group:
            out_rows.append([cat, mm, oldc, old, new])
        if len(group) > 1:
            out_rows.append([cat, "TOTAL", old_catalogs_total[cat],
                             sum(r[3] for r in group),
                             sum(r[4] for r in group)])

    out_rows.append(["ALL CATALOGS", "GRAND TOTAL", old_catalogs_grand,
                     sum(r[3] for r in rows), sum(r[4] for r in rows)])

    mapping_header = ["old_catalog_title", "new_catalog_title", "merge_methods",
                      "old_datasets"]
    mapping_rows = [list(r) for r in con.execute(MAPPING_SQL).fetchall()]

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(out_rows)

    write_xlsx([("catalog x merge_method", header, out_rows),
                ("old catalog mapping", mapping_header, mapping_rows)])

    print(f"wrote {len(out_rows)} rows -> {os.path.relpath(OUT)}")
    print(f"wrote {len(mapping_rows)} old catalogs -> "
          f"{os.path.relpath(OUT_XLSX)} (2 sheets)")
    print(f"old catalogs {old_catalogs_grand:,}, "
          f"old datasets {out_rows[-1][3]:,} -> new datasets {out_rows[-1][4]:,}")


def write_xlsx(sheets):
    """sheets: [(title, header, rows)] -> OUT_XLSX, header frozen and bold."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, name in enumerate(header, start=1):
            widest = max([len(name)] + [len(str(r[i - 1])) for r in rows])
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(i)].width = min(widest + 2, 60)
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    main()
