"""
Build dublin_catalog_lot2n3: the catalog-level Dublin Core table for lots 2 and 3.

Same columns, order and conventions as dublin_core_catalog_lot1
(build_dublin_core_catalog_lot1.py); source is remaining_catalog (12,150
data.gov.in catalogs, one row per nid) instead of catalog_metadata.

  - nid and Identifier[UUID] are the keys.
  - Title is the raw catalog title; new_catalog_title is the merged name
    (for the 50 catalogs with no dataset in either lot, the catalog's own
    title, 14 of them reworded by hand in remaining_catalog).
  - Description / Subject[sector] are the LLM-generated text
    (generated_description / generated_sector, llm_catalog_classifier_remaining.py);
    the raw body:value / field_sector:name only fill in if a generated value
    is missing.
  - Dates dd-mm-yyyy -> ISO YYYY-MM-DD. The 1970-01-01 "created" values are
    an export placeholder and become NULL.
  - Spatial Coverage: State/UT named in new_catalog_title (else Title); for
    State-jurisdiction catalogs naming none, the State/UT in the publisher
    ("Tamil Nadu,Planning..."); else "India" (lot 1 rule, spatial_from_title).
  - Creator / Contributor / Publisher: ministry/department, else the state
    department.
  - Subject[keyword]: top KEYWORD_SLOTS keywords of the datasets (lot 2 and
    lot 3 Subject[keyword], split on ; and ,) under the same new_catalog_title,
    ranked by how many catalogs use them. Catalogs with no dataset keywords
    (or no title) use their own raw keywords, truncated to the same limit.
  - Relation: data.gov.in landing pages of the catalog's datasets
    (Identifier[landing_page]), "; "-separated in nid order. Datasets point at
    their catalog by title; when several catalogs share a title the one with
    the nearest nid gets the dataset.
  - Relation[Is Part Of] = new_catalog_title. Has Part, Is Referenced By,
    Is Version Of, Has Version, Replaces, Is Replaced By: NULL (no verified
    source; lot 1's Replaces pairs were checked by hand, these have not been).

Creates a new table only. Usage:
  python build_dublin_catalog_lot2n3.py --dry-run
  python build_dublin_catalog_lot2n3.py [--replace] [--export]
  python build_dublin_catalog_lot2n3.py --export-only
"""

import argparse
import sys
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build_dublin_core_catalog_lot1 import spatial_from_title  # noqa: E402

DB_PATH = str(HERE.parent / "metadata.db")
SOURCE = "remaining_catalog"
DATASET_TABLES = ["dublin_core_remaining", "dublin_core_lot3"]
TABLE = "dublin_catalog_lot2n3"
REFERENCE = "dublin_core_catalog_lot1"
XLSX_PATH = HERE.parent.parent / "data" / "dublin_core_transformed" / f"{TABLE}.xlsx"
EXCEL_CELL_MAX = 32767
KEYWORD_SLOTS = 15

DATASETS_SQL = " UNION ALL ".join(
    f"""SELECT nid::BIGINT AS dataset_nid, trim("Relation[Catalog Title]") AS catalog_title,
               "Identifier[landing_page]" AS link, "Subject[keyword]" AS kw FROM {t}"""
    for t in DATASET_TABLES
)

# nid is VARCHAR in remaining_catalog; cast once.
CATALOGS_SQL = f"""
    SELECT *, nid::BIGINT AS cat_nid, trim(title) AS ttl,
           nullif(trim(new_catalog_title), '') AS nct
    FROM {SOURCE}
"""

DATASET_LINKS_SQL = f"""
    WITH ds AS ({DATASETS_SQL}), matched AS (
        SELECT ds.*, c.cat_nid,
               row_number() OVER (PARTITION BY ds.dataset_nid
                                  ORDER BY abs(c.cat_nid - ds.dataset_nid), c.cat_nid) AS rn
        FROM ds JOIN cats c ON c.ttl = ds.catalog_title
    )
    SELECT * EXCLUDE (rn) FROM matched WHERE rn = 1
"""

RELATIONS_SQL = """
    SELECT cat_nid, string_agg(link, '; ' ORDER BY dataset_nid) AS relation
    FROM dataset_links WHERE link IS NOT NULL AND link <> '' GROUP BY 1
"""

# Distinct keywords per (new title, original catalog), ranked by catalogs, then datasets.
KEYWORDS_SQL = f"""
    WITH kw AS (
        SELECT DISTINCT c.nct, c.cat_nid, d.dataset_nid, trim(k) AS keyword
        FROM dataset_links d JOIN cats c USING (cat_nid),
             unnest(regexp_split_to_array(d.kw, '[;,]')) AS s(k)
        WHERE c.nct IS NOT NULL AND trim(k) <> ''
    ), slots AS (
        SELECT nct, lower(keyword) AS lk, any_value(keyword) AS keyword,
               count(DISTINCT cat_nid) AS cats, count(DISTINCT dataset_nid) AS ds
        FROM kw GROUP BY 1, 2
    ), ranked AS (
        SELECT *, row_number() OVER (PARTITION BY nct ORDER BY cats DESC, ds DESC, lk) AS slot
        FROM slots
    )
    SELECT nct, string_agg(keyword, ', ' ORDER BY slot) AS keywords
    FROM ranked WHERE slot <= {KEYWORD_SLOTS} GROUP BY 1
"""

# Catalog's own keywords (';'-separated), first KEYWORD_SLOTS.
OWN_KEYWORDS_SQL = f"""
    SELECT cat_nid,
           array_to_string(list_slice(list_transform(
               list_filter(string_split(keywords, ';'), x -> trim(x) <> ''),
               x -> trim(x)), 1, {KEYWORD_SLOTS}), ', ') AS own_keywords
    FROM cats
"""


def iso(col: str) -> str:
    return f"""CASE WHEN c.{col} IN ('01-01-1970', '') THEN NULL
                    ELSE strftime(try_strptime(c.{col}, '%d-%m-%Y'), '%Y-%m-%d') END"""


SELECT_SQL = f"""
    SELECT
        c.cat_nid                                               AS "nid",
        c.uuid                                                  AS "Identifier[UUID]",
        c.ttl                                                   AS "Title",
        c.nct                                                   AS "new_catalog_title",
        coalesce(nullif(trim(c.generated_description), ''), c."body:value")
                                                                AS "Description",
        {iso('changed')}                                        AS "Date Modified",
        {iso('created')}                                        AS "Date Created",
        {iso('published_date')}                                 AS "Date Issued",
        CASE WHEN title_spatial(coalesce(c.nct, c.ttl)) <> 'India'
                  OR coalesce(nullif(c.field_ds_govt_type, ''), c.govt_type) <> 'State'
             THEN title_spatial(coalesce(c.nct, c.ttl))
             ELSE title_spatial(pub.p) END                      AS "Spatial Coverage",
        coalesce(nullif(c.field_ds_govt_type, ''), c.govt_type) AS "Jurisdiction",
        pub.p                                                   AS "Creator",
        pub.p                                                   AS "Contributor",
        pub.p                                                   AS "Publisher",
        coalesce(nullif(trim(c.generated_sector), ''), c."field_sector:name")
                                                                AS "Subject[sector]",
        coalesce(k.keywords, nullif(o.own_keywords, ''))        AS "Subject[keyword]",
        try_cast(c.from_api AS BIGINT)                          AS "Relation_from_api",
        try_cast(c.high_value_dataset AS BIGINT)                AS "High Value Dataset Category",
        try_cast(c.is_api_available AS BIGINT)                  AS "unmappable_is_api_available",
        try_cast(c.ogdp_view_count AS BIGINT)                   AS "unmappable_ogdp_view_count",
        'https://' || c.ogpl_module_domain_name || c.node_alias AS "Identifier",
        rel.relation                                            AS "Relation",
        'NDSAP'                                                 AS "Rights Statement",
        'Open'                                                  AS "Access Rights",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Has Part]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Is Referenced By]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Is Replaced By]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Replaces]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Is Version Of]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Has Version]",
        c.nct                                                   AS "Relation[Is Part Of]"
    FROM cats c
    CROSS JOIN LATERAL (SELECT coalesce(
        nullif(trim(c."field_ministry_department:name"), ''),
        nullif(trim(c."field_state_department:name"), ''),
        nullif(trim(c.state_department), '')) AS p) pub
    LEFT JOIN new_keywords k ON k.nct = c.nct
    LEFT JOIN own_keywords o ON o.cat_nid = c.cat_nid
    LEFT JOIN catalog_relations rel ON rel.cat_nid = c.cat_nid
    ORDER BY c.cat_nid
"""


def prepare(con: duckdb.DuckDBPyConnection) -> None:
    con.create_function("title_spatial", spatial_from_title, ["VARCHAR"], "VARCHAR")
    con.execute(f"CREATE OR REPLACE TEMP TABLE cats AS {CATALOGS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP TABLE dataset_links AS {DATASET_LINKS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP VIEW catalog_relations AS {RELATIONS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP TABLE new_keywords AS {KEYWORDS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP VIEW own_keywords AS {OWN_KEYWORDS_SQL}")


def report(con: duckdb.DuckDBPyConnection) -> None:
    n_ds = sum(con.execute(f"SELECT count(*) FROM {t}").fetchone()[0] for t in DATASET_TABLES)
    linked, cats = con.execute(
        "SELECT count(*), count(DISTINCT cat_nid) FROM dataset_links").fetchone()
    print(f"Datasets matched to a catalog: {linked} of {n_ds}, across {cats} catalogs")
    print(f"Catalogs without datasets: "
          f"{con.execute('SELECT count(*) FROM cats').fetchone()[0] - cats}")
    print(f"New titles with dataset keywords: "
          f"{con.execute('SELECT count(*) FROM new_keywords').fetchone()[0]}")
    if linked != n_ds:
        print(f"WARNING: {n_ds - linked} datasets matched no catalog title.")


def export_xlsx() -> None:
    con = duckdb.connect(DB_PATH, read_only=True)
    try:
        df = con.execute(f'SELECT * FROM {TABLE} ORDER BY new_catalog_title, "Title", nid').df()
    finally:
        con.close()
    links = df[["nid", "Title", "Relation"]].dropna(subset=["Relation"])
    links = (links.assign(Relation=links["Relation"].str.split("; "))
                  .explode("Relation")
                  .rename(columns={"nid": "catalog_nid", "Title": "catalog_title",
                                   "Relation": "dataset_link"}))
    relation = df["Relation"].fillna("")
    too_long = relation.str.len() > EXCEL_CELL_MAX
    counts = relation.str.count("; ") + 1
    df.loc[too_long, "Relation"] = (
        counts[too_long].astype(str) + " datasets; full list on sheet 'Relation'")
    longest = df.astype(str).apply(lambda s: s.str.len().max())
    if longest.max() > EXCEL_CELL_MAX:
        raise SystemExit(f"{longest.idxmax()} still exceeds the Excel cell limit.")
    # Excel caps a sheet at 1,048,576 rows; the link sheet is split if needed.
    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as xw:
        frames = [(TABLE, df)]
        step = 1_000_000
        for i in range(0, max(len(links), 1), step):
            frames.append(("Relation" if i == 0 else f"Relation_{i // step + 1}",
                           links.iloc[i:i + step]))
        for sheet, frame in frames:
            frame.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            widths = frame.astype(str).apply(lambda s: s.str.len().max())
            for i, col in enumerate(frame.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = (
                    min(60, max(len(col), int(widths[col])) + 2))
    print(f"Wrote {XLSX_PATH}: {len(df)} catalogs ({int(too_long.sum())} Relation cells "
          f"moved to sheet 'Relation'), {len(links)} dataset links.")


def main(dry_run: bool, replace: bool, export: bool, export_only: bool) -> None:
    if export_only:
        export_xlsx()
        return
    con = duckdb.connect(DB_PATH, read_only=dry_run)
    try:
        exists = con.execute(
            "SELECT count(*) FROM information_schema.tables WHERE table_name = ?", [TABLE]
        ).fetchone()[0]
        if exists and not replace and not dry_run:
            raise SystemExit(f"{TABLE} already exists; pass --replace to rebuild it.")
        prepare(con)
        report(con)
        if dry_run:
            print("[dry-run] No changes written.")
            return
        verb = "CREATE OR REPLACE TABLE" if replace else "CREATE TABLE"
        con.execute(f"{verb} {TABLE} AS {SELECT_SQL}")
        print(f"Created {TABLE}: {con.execute(f'SELECT count(*) FROM {TABLE}').fetchone()[0]} rows.")
        ref = [r[0] for r in con.execute(f"DESCRIBE {REFERENCE}").fetchall()]
        new = [r[0] for r in con.execute(f"DESCRIBE {TABLE}").fetchall()]
        if ref != new:
            raise SystemExit(f"Columns differ from {REFERENCE}: {ref} vs {new}")
        print(f"Columns match {REFERENCE} ({len(new)}).")
    finally:
        con.close()
    if export:
        export_xlsx()


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=f"Build {TABLE} from {SOURCE}.")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--replace", action="store_true")
    p.add_argument("--export", action="store_true")
    p.add_argument("--export-only", action="store_true")
    a = p.parse_args()
    main(a.dry_run, a.replace, a.export, a.export_only)
