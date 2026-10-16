#!/usr/bin/env python3
"""
Record counts per merge_method for lot 1 (dublin_core_metadata), lot 2
(dublin_core_remaining) and lot 3 (dublin_core_lot3), in the layout of "Record Count
for all datasets under Curate" (the lot 1 curate sheet built by
transformation/curate_collection_row_count.py). Lot 1 keeps a row_count column in its
table; lots 2 and 3 do not, so their files are counted here. Nothing is written to
the database.

Steps (the per-lot table, merge methods, download folder and output are in LOTS):

  uuids   (lot 2|3) Write the UUIDs of the datasets that get merged, so `count` can
          run on a machine without a current metadata.db.

  count   (lot 2|3) Count the data rows of those datasets' downloaded files
          ({uuid}.{ext} in --data-dir) and cache them in
          data/row_count_cache_<data-dir name>.csv, keyed by (name, size, mtime_ns);
          a rerun only opens new or changed files (the lot 2 drive is slow NTFS/FUSE
          on a 5400 rpm disk).

  import  (lot 2|3) Take the counts of a workbook this script exported on another
          machine (the one that holds the files) into the cache, for the datasets the
          cache has no row for. Only part of the lot 3 download is on this machine.

  export  (lot 1|2|3|all) One sheet per merge_method plus a Summary sheet. Only the
          sheets of the merged methods carry row_count and collection_total_row_count;
          the total is written once per collection and merged down its rows, as in
          the lot 1 sheet. `--lot all` writes the three workbooks and one more with
          the Summary of all lots side by side.

Row count = data rows after the header (a single header row is assumed, as lot 1's
`wc -l` minus 1 did), blank rows ignored:
  csv   records per the csv module, so quoted line breaks do not inflate the count.
        When one record runs over MAX_RECORD_LINES lines (an unbalanced quote
        swallowing the rows after it) the line count is used and noted.
  xls   first non-empty sheet via calamine; ".xls" files that are really HTML or
        delimited text are counted as such.
  json  length of the top-level list / "records" / "data".
Files that are not tables (pdf, kml, apk, ...) get no count and a note.

Usage:
  python3 merge_method_row_counts.py count --lot 2 [--db PATH] [--workers 6] [--recount NOTE_PREFIX]
  python3 merge_method_row_counts.py export --lot 2 [--db PATH] [--out PATH]

  # lot 3, counting on another machine that holds the files:
  python3 merge_method_row_counts.py uuids --lot 3 --out lot3_merged_uuids.txt
  python3 merge_method_row_counts.py count --lot 3 --uuids-file lot3_merged_uuids.txt \
          --data-dir /path/to/lot3_downloads
  python3 merge_method_row_counts.py export --lot 3 --cache data/row_count_cache_lot3_downloads.csv
  # ... or, with only the workbook exported there:
  python3 merge_method_row_counts.py import --lot 3 --workbook "Record Count for Lot 3 datasets by merge method.xlsx"

  python3 merge_method_row_counts.py export --lot all
"""

import argparse
import csv
import json
import os
import re
import sys
import time
from multiprocessing import Pool
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
DB_PATH = HERE.parent / "metadata.db"
CACHE_FIELDS = ["name", "size", "mtime_ns", "row_count", "kind", "note"]

# Left out on purpose (lot 2 only): 'failed to download' (4,847 datasets on 2026-10-09, none
# downloaded; the former 'no header' rows are among them). The Summary notes give the live count.
MERGE_METHODS = ["curate", "direct", "change_in_header", "non-mergeable", "to be removed"]
MERGED = ["curate", "direct", "change_in_header"]   # members roll up into a Collection

# Where each lot's columns come from. None = the lot does not store it (blank / NA).
# row_count: the table's column (lot 1), or None = counted from the files in data_dir.
DIMS = ["Year", "Month", "Financial_Year", "Quarter", "State"]
LOTS = {
    "1": {
        "table": "dublin_core_metadata",
        "data_dir": None,
        "row_count": "row_count",
        "new_title": "new_title",
        "new_catalog_title": "new_catalog_title",
        "dims": {"Year": "Year", "Month": "Month", "Financial_Year": "Financial_Year",
                 "Quarter": "Quarter", "State": None},
        # lot 1 has two header-change methods where lots 2 and 3 have change_in_header
        "methods": ["curate", "direct", "Minor change in header", "major_changes_in_headers",
                    "non-mergeable", "to be removed"],
        "merged": ["curate", "direct", "Minor change in header", "major_changes_in_headers"],
    },
    "2": {
        "table": "dublin_core_remaining",
        "data_dir": REPO / "data" / "final_batch_downloads",
        "row_count": None,
        "new_title": "new_title",
        "new_catalog_title": "new_catalog_title",
        "dims": {"Year": "year", "Month": "month", "Financial_Year": "financial_year",
                 "Quarter": "quarter", "State": "state"},
        "methods": MERGE_METHODS,
        "merged": MERGED,
    },
    "3": {
        "table": "dublin_core_lot3",
        "data_dir": REPO / "data" / "lot3_downloads",
        "row_count": None,
        "new_title": "new_title",        # Collection, else Title (added 2026-10-09)
        "new_catalog_title": "new_catalog_title",   # added 2026-10-07
        "dims": {"Year": None, "Month": None, "Financial_Year": None, "Quarter": None,
                 "State": "spatial_states"},
        "methods": MERGE_METHODS,
        "merged": MERGED,
    },
}
FILE_LOTS = sorted(k for k, v in LOTS.items() if v["row_count"] is None)
IMPORTED = "imported"   # cache kind of a count taken from another machine's workbook

CSV_FIELD_LIMIT = 10_000_000   # an unbalanced quote hits this instead of eating RAM
MAX_RECORD_LINES = 50
NOT_TABULAR = {"pdf", "vnd.google-earth.kml+xml", "vnd.android.package-archive",
               "css", "html"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0"
ZIP_MAGIC = b"PK\x03\x04"


# --------------------------------------------------------------------------- count

def _blank(line: str) -> bool:
    return not line.strip(" \t\r\n,;\"'")


def count_delimited(path: Path, encoding: str):
    """-> (rows, note). Streams, so a 1.3 GB csv costs no memory."""
    lines = 0

    def feed(f):
        nonlocal lines
        for line in f:
            if not _blank(line):
                lines += 1
            yield line

    records = 0
    longest = 0                    # most lines taken by one record
    broke = None
    with open(path, encoding=encoding, errors="replace", newline="") as f:
        it = feed(f)
        try:
            before = 0
            for row in csv.reader(it):
                longest = max(longest, lines - before)
                before = lines
                if any(c.strip() for c in row):
                    records += 1
        except csv.Error as e:
            broke = str(e)
            for _ in it:           # finish the line count past the bad field
                pass

    # Quoted line breaks (multi-line headers are common) are fine; a record running
    # over MAX_RECORD_LINES lines is an unbalanced quote eating the rows after it.
    if broke or longest > MAX_RECORD_LINES:
        return max(lines - 1, 0), "unbalanced quotes: line count used"
    return max(records - 1, 0), ""


def count_excel(path: Path):
    from python_calamine import CalamineWorkbook
    with open(path, "rb") as f:    # by content: some ".xls" files are xlsx inside
        wb = CalamineWorkbook.from_filelike(f)
    names = wb.sheet_names
    for i in range(len(names)):
        rows = wb.get_sheet_by_index(i).to_python(skip_empty_area=True)
        n = sum(1 for r in rows if any(str(c).strip() for c in r if c is not None))
        if n:
            note = f"first of {len(names)} sheets ({names[i]})" if len(names) > 1 else ""
            return max(n - 1, 0), note
    return 0, "no non-empty sheet"


def count_json(path: Path):
    try:
        with open(path, "rb") as f:
            doc = json.load(f)
    except ValueError:
        if path.stat().st_size == 262_144:     # 372 lot 3 files stop at exactly 256 KiB
            return None, "not parseable as json: exactly 256 KiB, download likely truncated"
        return None, "not parseable as json"
    if isinstance(doc, list):
        return len(doc), ""
    if isinstance(doc, dict):
        for key in ("records", "data", "rows", "features"):
            if isinstance(doc.get(key), list):
                return len(doc[key]), ""
        lists = [k for k, v in doc.items() if isinstance(v, list)]
        if len(lists) == 1:                 # {"MMI Irrigation Projects": [...]}
            return len(doc[lists[0]]), ""
    return None, "unrecognised json layout"


def count_html(path: Path):
    with open(path, encoding="latin-1") as f:
        text = f.read().lower()
    trs = text.count("<tr")
    return max(trs - 1, 0), "html table"


def count_file(entry):
    """entry = (path, size, mtime_ns) -> cache row dict. Never raises."""
    path, size, mtime_ns = entry
    path = Path(path)
    ext = path.name.partition(".")[2].lower()
    out = {"name": path.name, "size": size, "mtime_ns": mtime_ns,
           "row_count": "", "kind": ext, "note": ""}
    try:
        if size == 0:
            out["note"] = "empty file"
            return out
        if ext in NOT_TABULAR:
            out["note"] = f"not tabular ({ext})"
            return out
        with open(path, "rb") as f:
            head = f.read(4096)
        if ext == "json":
            n, note = count_json(path)
        elif head.startswith(OLE_MAGIC) or (head.startswith(ZIP_MAGIC) and ext in ("xls", "xlsx", "ods")):
            out["kind"] = "excel"
            n, note = count_excel(path)
        elif b"<table" in head.lower() or head.lstrip().lower().startswith((b"<html", b"<!doctype")):
            out["kind"] = "html"
            n, note = count_html(path)
        else:
            enc = "utf-16" if head.startswith((b"\xff\xfe", b"\xfe\xff")) else "latin-1"
            if ext != "csv":
                out["kind"] = f"{ext} (text)"
            n, note = count_delimited(path, enc)
        out["row_count"] = "" if n is None else n
        out["note"] = note
    except Exception as e:  # noqa: BLE001 - one bad file must not stop a 1 h pass
        out["note"] = f"error: {type(e).__name__}: {str(e)[:150]}"
    return out


def lot_paths(args, key):
    """-> (download folder, row count cache) of a lot; --data-dir / --cache override."""
    data_dir = Path(getattr(args, "data_dir", None) or LOTS[key]["data_dir"])
    cache = getattr(args, "cache", None)
    return data_dir, Path(cache) if cache else REPO / "data" / f"row_count_cache_{data_dir.name}.csv"


def load_cache(path):
    cache = {}
    if path.exists():
        with open(path, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r.get("note") is None:  # line cut short while a count is still writing
                    continue
                cache[r["name"]] = r       # later lines win
    return cache


def merged_uuids(db, lot):
    import duckdb
    conn = duckdb.connect(db, read_only=True)
    try:
        return {u for (u,) in conn.execute(
            f'SELECT "Identifier[UUID]" FROM {lot["table"]} '
            f'WHERE merge_method IN ({", ".join("?" * len(lot["merged"]))})',
            lot["merged"]).fetchall() if u}
    finally:
        conn.close()


def cmd_uuids(args):
    lot = LOTS[args.lot]
    wanted = sorted(merged_uuids(args.db, lot))
    Path(args.out).write_text("\n".join(wanted) + "\n")
    print(f"{len(wanted):,} UUIDs (merge_method in {lot['merged']}) -> {args.out}")


def cmd_count(args):
    csv.field_size_limit(CSV_FIELD_LIMIT)
    t0 = time.time()
    lot = LOTS[args.lot]
    data_dir, path = lot_paths(args, args.lot)
    if args.uuids_file:
        wanted = set(Path(args.uuids_file).read_text().split())
    else:
        wanted = merged_uuids(args.db, lot)
    print(f"{len(wanted):,} datasets with merge_method in {lot['merged']}", flush=True)
    print(f"Listing {data_dir} ...", flush=True)
    entries = []
    with os.scandir(data_dir) as it:
        for e in it:
            if e.is_file() and e.name.partition(".")[0] in wanted:
                st = e.stat()
                entries.append((e.inode(), e.path, st.st_size, st.st_mtime_ns))
    entries.sort()                           # MFT order: fewer seeks on this disk
    print(f"  {len(entries):,} of their files on disk ({time.time() - t0:.0f}s)", flush=True)

    cache = load_cache(path)
    retry = ("error:",) + tuple(args.recount)

    def fresh(c, s, m):
        return (c and c["size"] == str(s) and c["mtime_ns"] == str(m)
                and not c["note"].startswith(retry))

    todo = [(p, s, m) for _, p, s, m in entries
            if not fresh(cache.get(os.path.basename(p)), s, m)]
    print(f"  cached: {len(entries) - len(todo):,}   to count: {len(todo):,}", flush=True)
    if args.limit:
        todo = todo[: args.limit]
    if not todo:
        return

    new_file = not path.exists()
    t1 = time.time()
    done = 0
    with open(path, "a", newline="", encoding="utf-8") as f, Pool(args.workers) as pool:
        w = csv.DictWriter(f, fieldnames=CACHE_FIELDS)
        if new_file:
            w.writeheader()
        for row in pool.imap_unordered(count_file, todo, chunksize=8):
            w.writerow(row)
            done += 1
            if done % 2000 == 0 or done == len(todo):
                f.flush()
                rate = done / (time.time() - t1)
                eta = (len(todo) - done) / rate if rate else 0
                print(f"  {done:,}/{len(todo):,}  {rate:.0f} files/s  eta {eta / 60:.0f} min",
                      flush=True)
    print(f"Done in {(time.time() - t0) / 60:.1f} min -> {path}", flush=True)


def cmd_import(args):
    """Counts from a workbook this script exported where the files are -> the cache.
    Only for datasets the cache has no row for; a dataset counted in both must agree."""
    import duckdb
    from openpyxl import load_workbook

    lot = LOTS[args.lot]
    _, path = lot_paths(args, args.lot)
    conn = duckdb.connect(args.db, read_only=True)
    try:
        table = {_nid(nid): (uuid, _val(title)) for nid, uuid, title in conn.execute(
            f'SELECT nid, "Identifier[UUID]", "Title" FROM {lot["table"]} '
            f'WHERE merge_method IN ({", ".join("?" * len(lot["merged"]))})', lot["merged"]).fetchall()}
    finally:
        conn.close()
    have = load_counts(path)
    wb = load_workbook(args.workbook, read_only=True)
    add, same, differ, unknown, no_file = [], 0, [], 0, 0
    for method in lot["merged"]:
        if sheet_title(method) not in wb.sheetnames:
            continue
        rows = wb[sheet_title(method)].iter_rows(values_only=True)
        head = list(next(rows))
        i_nid, i_title, i_n, i_note = (head.index(h) for h in
                                       ("nid", "Old Dataset Title", "row_count", "row_count_note"))
        for r in rows:
            r = list(r) + [None] * (len(head) - len(r))     # trailing empty cells are not stored
            uuid, title = table.get(_nid(r[i_nid]), (None, None))
            n, note = _count(r[i_n]), str(r[i_note] or "")
            if uuid is None or title != _val(r[i_title]):
                unknown += 1          # not a merged dataset here (any more), or under another title
            elif n is None and note == "file not downloaded":
                no_file += 1          # no file there either: nothing to take over
            elif uuid in have:
                if have[uuid]["row_count"] == ("" if n is None else str(n)):
                    same += 1
                else:
                    differ.append((r[i_nid], have[uuid]["row_count"], n))
            else:
                add.append({"name": uuid, "size": "", "mtime_ns": "", "kind": IMPORTED,
                            "row_count": "" if n is None else n, "note": note})
    print(f"{args.workbook}: {len(add):,} counts to take over, {same:,} already counted here with the "
          f"same result, {no_file:,} without a file there, {unknown:,} rows not matched to "
          f"{lot['table']}")
    if differ:
        sys.exit(f"{len(differ):,} datasets counted here have another count in the workbook, "
                 f"e.g. (nid, here, workbook) {differ[:5]} - nothing imported")
    if not add:
        return
    new_file = not path.exists()
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CACHE_FIELDS)
        if new_file:
            w.writeheader()
        w.writerows(add)
    print(f"  added to {path}")


# -------------------------------------------------------------------------- export

# (header, width) in sheet order; widths follow the lot 1 curate sheet.
COLUMNS = [
    ("nid", 10), ("Old Dataset Title", 52.5), ("New Dataset Title", 52.5),
    ("Old Catalog title", 43.75), ("New Catalog Title", 35), ("merge_add_columns", 19.25),
    ("Year", 7), ("Month", 8.75), ("Financial_Year", 13.13), ("Quarter", 7.88),
    ("State", 16), ("District", 16), ("row_count", 10.5),
    ("collection_total_row_count", 22.75), ("row_count_note", 30),
]
COUNT_COLUMNS = {"row_count", "collection_total_row_count", "row_count_note"}  # MERGED sheets only
NA_FILL = "NA"   # dimension columns: a visible NA, as in the lot 1 sheet
NA_VALUES = {"", "<na>", "na", "nan", "none", "null"}
# "... for Srikakulam District of Andhra Pradesh, 1991" (lot 2 Census),
# "... in the Dharmapuri District : DHB 2017-18" (lot 3 District Hand Books)
DISTRICT_RE = re.compile(r".*\b(?:for|in|of)\s+(?:the\s+)?(?P<d>[A-Z][^:,;]*?)\s+District"
                         r"(?=\s+of\b|\s*[:,;(]|\s*$|\s+\d)")

SUMMARY_HEADERS = [
    ("merge_method", 18), ("old datasets", 13), ("new datasets", 13), ("old catalogs", 13),
    ("new catalogs", 13), ("datasets with row_count", 14), ("datasets without row_count", 14),
    ("total rows", 14), ("collections missing a member's count", 18),
]


def _val(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s.lower() in NA_VALUES else s


def _int_or_text(v):
    v = _val(v)
    return int(float(v)) if v is not None and re.fullmatch(r"\d+(\.0)?", v) else v


def _count(v):
    """A stored or exported row count -> int, None when there is none."""
    v = _val(v)
    try:
        return None if v is None else int(float(v))
    except ValueError:
        return None


def _nid(v):
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def sheet_title(method):
    title = f"{method} datasets"
    return title if len(title) <= 31 else method     # Excel caps sheet names at 31 characters


def load_counts(path):
    """uuid -> cache row, from the count (and import) step."""
    return {name.partition(".")[0]: r for name, r in load_cache(path).items()}


def _col(name):
    return f'"{name}"' if name else "NULL"


def build_rows(conn, lot, method, counts, on_disk):
    merged = method in lot["merged"]
    q = f"""
        SELECT nid, "Title", "Collection", {_col(lot["new_title"])}, "Relation[Catalog Title]",
               {_col(lot["new_catalog_title"])}, merge_add_columns,
               {", ".join(_col(lot["dims"][d]) for d in DIMS)}, "Identifier[UUID]",
               {_col(lot["row_count"])}
        FROM {lot["table"]} WHERE merge_method = ?
    """
    out = []
    for (nid, title, coll, new_title, old_cat, new_cat, mac, year, month, fy,
         quarter, state, uuid, stored) in conn.execute(q, [method]).fetchall():
        coll, new_title, mac = _val(coll), _val(new_title), _val(mac)
        new_ds = coll or new_title
        if new_ds is None and method == "non-mergeable":
            new_ds = _val(title)       # blank new_title: an unmerged dataset keeps its title
        n, note = None, None
        if merged and lot["row_count"]:
            n = _count(stored)
            note = None if n is not None else "no row_count stored"
        elif merged:
            c = counts.get(uuid)
            if c is None:
                note = "not counted yet (run count)" if uuid in on_disk else "file not downloaded"
            else:
                n = int(c["row_count"]) if c["row_count"] != "" else None
                note = c["note"]
        district = None
        if mac and "district" in mac.lower():
            m = DISTRICT_RE.match(title or "")
            district = m.group("d").strip() if m else None
        out.append({
            "nid": int(_nid(nid)) if _nid(nid).isdigit() else nid,
            "Old Dataset Title": _val(title),
            "New Dataset Title": new_ds,
            "Old Catalog title": _val(old_cat),
            "New Catalog Title": _val(new_cat),
            "merge_add_columns": mac,
            "Year": _int_or_text(year),
            "Month": _val(month),
            "Financial_Year": _val(fy),
            "Quarter": _int_or_text(quarter),
            "State": _val(state),
            "District": district,
            "row_count": n,
            "row_count_note": note or None,
            # grouping: a Collection in the merged methods, otherwise the dataset itself
            "_group": coll if (merged and coll) else f"nid:{nid}",
        })
    out.sort(key=lambda r: ((r["New Dataset Title"] or "").lower(), r["_group"],
                            r["nid"] if isinstance(r["nid"], int) else 0))
    return out


def write_sheet(wb, method, rows, merged):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange

    cols = [(h, w) for h, w in COLUMNS if merged or h not in COUNT_COLUMNS]
    names = [h for h, _ in cols]
    total_col = names.index("collection_total_row_count") + 1 if "collection_total_row_count" in names else None

    ws = wb.create_sheet(sheet_title(method))
    ws.freeze_panes = "A2"
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    head_font, head_align = Font(bold=True), Alignment(vertical="center", wrap_text=True)
    header = []
    for h in names:
        cell = WriteOnlyCell(ws, value=h)
        cell.font, cell.alignment = head_font, head_align
        header.append(cell)
    ws.append(header)

    # collection totals over the members that have a count
    totals, missing = {}, {}
    for r in rows:
        g = r["_group"]
        if r["row_count"] is None:
            missing[g] = True
        else:
            totals[g] = totals.get(g, 0) + r["row_count"]

    center = Alignment(horizontal="center", vertical="center")
    excel_row, start, current = 1, None, object()

    def close(end):
        if total_col and start is not None and end > start:
            ws.merged_cells.ranges.add(CellRange(min_col=total_col, min_row=start,
                                                 max_col=total_col, max_row=end))

    for r in rows:
        excel_row += 1
        values = []
        for h in names:
            v = r.get(h)
            if h in ("Year", "Month", "Financial_Year", "Quarter", "State", "District") and v is None:
                v = NA_FILL
            values.append(v)
        if total_col:
            if r["_group"] != current:
                close(excel_row - 1)
                current, start = r["_group"], excel_row
                cell = WriteOnlyCell(ws, value=totals.get(current))
                cell.alignment = center
                values[total_col - 1] = cell
            else:
                values[total_col - 1] = None
        ws.append(values)
    close(excel_row)

    groups = {r["_group"] for r in rows}
    return {
        "groups": len(groups),
        "groups_missing": sum(1 for g in groups if g in missing and not g.startswith("nid:")),
    }


def summary_row(method, rows, merged):
    counted = [r["row_count"] for r in rows if r["row_count"] is not None]
    # a collection is one new dataset, an unmerged dataset is its own; removed ones are none
    keep = method != "to be removed"
    colls = {r["_group"] for r in rows if keep and not r["_group"].startswith("nid:")}
    standalone = sum(1 for r in rows if keep and r["_group"].startswith("nid:"))
    return {
        "merge_method": method,
        "old datasets": len(rows),
        "new datasets": len(colls) + standalone,
        "old catalogs": len({r["Old Catalog title"] for r in rows if r["Old Catalog title"]}),
        "new catalogs": len({r["New Catalog Title"] for r in rows if r["New Catalog Title"]}),
        "datasets with row_count": len(counted) if merged else None,
        "datasets without row_count": len(rows) - len(counted) if merged else None,
        "total rows": sum(counted) if merged else None,
        "_colls": colls,
        "_standalone": standalone,
        "_old_cats": {r["Old Catalog title"] for r in rows if r["Old Catalog title"]},
        "_new_cats": {r["New Catalog Title"] for r in rows if r["New Catalog Title"]},
    }


def summary_total(summary, label="TOTAL"):
    """Sums, except what one collection or catalog can add to twice: those are counted over all rows."""
    total = {h: sum(s[h] or 0 for s in summary) for h, _ in SUMMARY_HEADERS[1:]}
    for h, key in (("new datasets", "_colls"), ("old catalogs", "_old_cats"), ("new catalogs", "_new_cats")):
        total[key] = set().union(*(s[key] for s in summary))
        total[h] = len(total[key])
    total["_standalone"] = sum(s["_standalone"] for s in summary)
    total["new datasets"] += total["_standalone"]
    total["merge_method"] = label
    return total


# {merged} = the lot's merged methods, {left_out} = the methods it has that get no sheet
SOURCE_NOTE = "Source: {table} (lot {lot}), one sheet per merge_method.{left_out}"
MERGED_NOTE = ("Row counts are only for the datasets that get merged ({merged}); the non-mergeable and "
               "to be removed sheets list the datasets without counts.")
FILE_COUNT_NOTE = ("row_count = data rows in the downloaded file after one header row, blank rows "
                   "skipped. csv: records as parsed, so quoted line breaks do not add rows; xls: first "
                   "non-empty sheet; json: record count. Files that are not tables, or were never "
                   "downloaded, have no count - see row_count_note.")
TOTAL_NOTES = [
    "collection_total_row_count = sum of row_count over the collection's datasets that have a count; "
    "it is merged down the collection's rows.",
    "new datasets, old catalogs and new catalogs are distinct counts, so TOTAL is counted over all "
    "rows, not summed: a collection with members under two merge methods is one new dataset.",
]
SUMMARY_NOTES = {
    "1": [
        SOURCE_NOTE,
        "New Dataset Title = Collection; datasets without one (non-mergeable, to be removed) show "
        "new_title.",
        "Year, Month, Financial_Year, Quarter: the stored columns (NA = empty). Lot 1 stores no state, "
        "so State is NA. District: taken from the title ('... for X District of Y') where "
        "merge_add_columns names district. Other merge_add_columns dimensions (facility, ...) are "
        "not shown.",
        MERGED_NOTE,
        "row_count = the row_count column of dublin_core_metadata (lines of the downloaded file minus "
        "one header line). Datasets where it is empty have no count - see row_count_note.",
    ],
    "2": [
        SOURCE_NOTE,
        "New Dataset Title = Collection; datasets without one (non-mergeable, to be removed) show "
        "new_title.",
        "Year, Month, Financial_Year, Quarter, State: the stored columns (NA = empty). District: taken "
        "from the title ('... for X District of Y') where merge_add_columns names district. Other "
        "merge_add_columns dimensions (airline, jail_type, ...) have no stored value.",
        MERGED_NOTE,
        FILE_COUNT_NOTE,
    ],
    "3": [
        SOURCE_NOTE,
        "New Dataset Title = Collection; datasets without one (non-mergeable, to be removed) show "
        "new_title, which is their Title.",
        "Lot 3 stores no Year, Month, Financial_Year or Quarter, so those columns are NA. State = "
        "spatial_states. District: taken from the title ('... for X District of Y') where "
        "merge_add_columns names district. Other merge_add_columns dimensions (date, crop, season, "
        "...) have no stored value.",
        MERGED_NOTE,
        FILE_COUNT_NOTE,
    ],
}
# Why a lot's other merge methods get no sheet.
LEFT_OUT_WHY = {"failed to download": "none of their files could be downloaded; the former 'no header' "
                                      "datasets are among them"}


def left_out_text(left_out):
    if not left_out:
        return ""
    parts = []
    for m, n in sorted(left_out.items(), key=lambda kv: -kv[1]):
        why = LEFT_OUT_WHY.get(m)
        parts.append(f"'{m}' ({n:,} datasets{'; ' + why if why else ''})")
    return " Left out: " + ", ".join(parts) + "."


def lot_notes(key, left_out, counts):
    lot = LOTS[key]
    notes = [n.format(table=lot["table"], lot=key, merged=", ".join(lot["merged"]),
                      left_out=left_out_text(left_out)) for n in SUMMARY_NOTES[key]]
    imported = sum(1 for c in counts.values() if c["kind"] == IMPORTED)
    if imported:
        notes.append(f"Only part of the lot {key} download is on the machine this workbook was made on: "
                     f"{len(counts) - imported:,} files were counted there, the other {imported:,} were "
                     f"counted the same way on the machine that holds all the files. Every file counted "
                     f"on both has the same count.")
    return notes + TOTAL_NOTES


def _bold_row(ws, values):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    row = [WriteOnlyCell(ws, value=v) for v in values]
    for c in row:
        c.font = Font(bold=True)
    return row


def write_summary(ws, blocks, notes):
    """blocks = [(label, summary rows)]. One lot: the reference layout. Several: a leading lot
    column, each lot's rows and total, then the total of all lots."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    many = len(blocks) > 1
    headers = ([("lot", 8)] if many else []) + SUMMARY_HEADERS
    for i, (_, w) in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    bold, wrap = Font(bold=True), Alignment(vertical="center", wrap_text=True)
    head = []
    for h, _ in headers:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.alignment = bold, wrap
        head.append(c)
    ws.append(head)
    totals = []
    for label, summary in blocks:
        lead = [label] if many else []
        for s in summary:
            ws.append(lead + [s[h] for h, _ in SUMMARY_HEADERS])
        totals.append(summary_total(summary))
        ws.append(_bold_row(ws, lead + [totals[-1][h] for h, _ in SUMMARY_HEADERS]))
        if many:
            ws.append([])
    if many:
        # a collection belongs to one lot, so new datasets add up; a catalog can hold datasets of two
        grand = summary_total(totals)
        grand["new datasets"] = sum(t["new datasets"] for t in totals)
        ws.append(_bold_row(ws, ["All lots"] + [grand[h] for h, _ in SUMMARY_HEADERS]))
    ws.append([])
    ws.append([WriteOnlyCell(ws, value="Notes")])
    for n in notes:
        ws.append([n])


def export_lot(args, key):
    """Write one lot's workbook -> (its summary rows, the methods left out)."""
    import duckdb
    from openpyxl import Workbook

    t0 = time.time()
    lot = LOTS[key]
    counts, on_disk = {}, set()
    if lot["row_count"] is None:
        data_dir, path = lot_paths(args, key)
        counts = load_counts(path)
        if not counts:
            raise SystemExit(f"No row counts in {path} - run the count step first.")
        if data_dir.is_dir():
            on_disk = {name.partition(".")[0] for name in os.listdir(data_dir)}
            uncounted = len((on_disk & merged_uuids(args.db, lot)) - counts.keys())
            if uncounted:
                print(f"  WARNING: {uncounted:,} downloaded files have no count yet - run the count step",
                      flush=True)
        else:
            # counted elsewhere (e.g. lot 3 on EC2): anything missing from the cache had no file there
            print(f"  {data_dir} not here; datasets missing from {path.name} are 'file not downloaded'",
                  flush=True)
    conn = duckdb.connect(args.db, read_only=True)
    try:
        per_method = {m: build_rows(conn, lot, m, counts, on_disk) for m in lot["methods"]}
        left_out = {m: n for m, n in conn.execute(
            f"SELECT merge_method, count(*) FROM {lot['table']} GROUP BY 1").fetchall()
            if m not in lot["methods"]}
    finally:
        conn.close()
    if left_out:
        print(f"  not exported (not in the lot's methods): {left_out}", flush=True)

    wb = Workbook(write_only=True)
    ws_sum = wb.create_sheet("Summary")          # first tab; filled once totals are known
    summary = []
    for m in lot["methods"]:
        if not per_method[m]:
            continue
        merged = m in lot["merged"]
        info = write_sheet(wb, m, per_method[m], merged)
        s = summary_row(m, per_method[m], merged)
        s["collections missing a member's count"] = info["groups_missing"] if merged else None
        summary.append(s)
        print(f"  {m:<24} {len(per_method[m]):>7,} datasets  {info['groups']:>7,} groups", flush=True)
    write_summary(ws_sum, [(f"Lot {key}", summary)], lot_notes(key, left_out, counts))

    out = Path(getattr(args, "out", None)
               or REPO / "data" / f"Record Count for Lot {key} datasets by merge method.xlsx")
    wb.save(out)
    print(f"Wrote {out} in {time.time() - t0:.0f}s")
    return summary, left_out


def cmd_export(args):
    from openpyxl import Workbook

    if args.lot != "all":
        export_lot(args, args.lot)
        return
    if args.out or args.cache or args.data_dir:
        sys.exit("--out, --cache and --data-dir go with one lot, not with --lot all")
    blocks, notes = [], []
    for key in sorted(LOTS):
        print(f"Lot {key} ({LOTS[key]['table']})", flush=True)
        summary, left_out = export_lot(args, key)
        blocks.append((f"Lot {key}", summary))
        notes.append(f"Lot {key} = {LOTS[key]['table']}; merged methods: {', '.join(LOTS[key]['merged'])}."
                     + left_out_text(left_out))
    notes += [
        "Each block is the Summary sheet of 'Record Count for Lot N datasets by merge method.xlsx', "
        "which lists the datasets behind it and explains its columns.",
        "Row counts are only for the datasets that get merged. Lot 1: the row_count column of its "
        "table. Lots 2 and 3: data rows counted in the downloaded files.",
        "new datasets, old catalogs and new catalogs are distinct counts, so a lot's TOTAL is counted "
        "over all its rows, not summed. All lots: new datasets is the sum of the three lots; the "
        "catalogs are counted over all lots, as one catalog can hold datasets of two lots.",
    ]
    wb = Workbook(write_only=True)
    write_summary(wb.create_sheet("Summary"), blocks, notes)
    out = REPO / "data" / "Record Count for all Lots by merge method - Summary.xlsx"
    wb.save(out)
    print(f"Wrote {out}")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def lot_args(p, lots, files=True):
        p.add_argument("--lot", choices=lots, required=True)
        p.add_argument("--db", default=str(DB_PATH))
        if files:
            p.add_argument("--data-dir", default=None,
                           help="folder of {uuid}.{ext} downloads (default: the lot's, see LOTS)")
            p.add_argument("--cache", default=None,
                           help="row count cache (default: data/row_count_cache_<data-dir name>.csv)")

    u = sub.add_parser("uuids", help="write the merged datasets' UUIDs, one per line")
    lot_args(u, FILE_LOTS, files=False)
    u.add_argument("--out", required=True)
    c = sub.add_parser("count", help="count rows of the merged datasets' files (cached)")
    lot_args(c, FILE_LOTS)
    c.add_argument("--uuids-file", default=None,
                   help="UUIDs to count (from `uuids`); skips the database")
    c.add_argument("--workers", type=int, default=6)
    c.add_argument("--limit", type=int, default=None, help="only count N uncached files (testing)")
    c.add_argument("--recount", action="append", default=[], metavar="NOTE_PREFIX",
                   help="also recount cached files whose note starts with this "
                        "(files with 'error:' notes are always retried)")
    i = sub.add_parser("import", help="take the counts of a workbook exported where the files are")
    lot_args(i, FILE_LOTS)
    i.add_argument("--workbook", required=True)
    e = sub.add_parser("export", help="write the per-merge_method workbook")
    lot_args(e, sorted(LOTS) + ["all"])
    e.add_argument("--out", default=None,
                   help='default: data/Record Count for Lot <N> datasets by merge method.xlsx')
    args = ap.parse_args()
    {"uuids": cmd_uuids, "count": cmd_count, "import": cmd_import, "export": cmd_export}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
