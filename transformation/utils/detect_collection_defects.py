#!/usr/bin/env python3
"""
Flag wrong collections among merge_method='curate' rows of dublin_core_remaining.

A collection is a set of datasets that will be vertically merged into one
published dataset. The curate collections were built by grouping on a
normalised title base plus an exact header signature; where one title base had
several header signatures the group was split and each part kept THE PARENT'S
NAME with a " (schema N)" suffix appended. The name was never regenerated, so
it describes one arbitrary member rather than the group.

Example: "Block wise Distribution of Surface Flow Schemes (Other Storage) ...
(schema 2)" is four separate collections sharing that one name. The grouping is
right (schema 1 = Village level, 2 = Block, 3 = District, 4 = State), but three
of the four names claim "Block wise" and all four claim "Other Storage" when
each collection actually spans 2 scheme types x 4 storage types -- a
distinction that lives only in the title, not in the headers, and that
merge_add_columns (NULL for all four) would lose on merge.

This script writes the evidence needed to decide each case, it does not decide.
It is READ-ONLY: the DB is opened read_only=True, always.

Flags (a collection can carry several):
  NAME_PINS_VALUE   a token in the name appears in < --name-threshold of the
                    member titles, i.e. the name pins one member's value
  SCHEMA_N_SUFFIX   the name ends in " (schema N)" -- meaningless to a reader,
                    needs a real discriminator
  HEADER_NULL_MIX   members with and without extracted headers sit together, so
                    the header split could not separate them
  MULTI_HEADER_SIG  more than one distinct non-null header signature (Rule A
                    violation -- these members cannot be stacked at all)
  MAC_REDUNDANT     merge_add_columns names a dimension that is already a
                    header column, so the merge would add a duplicate
  MAC_MISSING_DIM   a title names a dimension the headers lack and
                    merge_add_columns does not list -- it is lost on merge

Output: an xlsx (and a csv beside it) with one row per flagged collection,
carrying the evidence plus five blank columns to fill in by hand:
new_collection_name, new_merge_add_columns, action, split_rule, notes.
Feed the filled workbook to apply_collection_review.py.

Decisions live in collection_decisions.csv, not in the workbook, and are
overlaid onto it on every run. The detector can therefore be re-run freely --
after headers get filled in, after a partial apply -- without losing hand work.

Usage:
  python3 detect_collection_defects.py                    # report + workbook
  python3 detect_collection_defects.py --summary-only     # counts, no workbook
  python3 detect_collection_defects.py --merge-method change_in_header
"""

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import duckdb

HERE = Path(__file__).resolve().parent
# Order matters: transformation/ also has a dataset_headers.py -- a different
# module that connects to a hardcoded (and on this machine non-existent) DB path
# at import time. utils/ must win, so it is inserted last.
sys.path.insert(0, str(HERE.parent))     # dataset_merge
sys.path.insert(0, str(HERE))            # sibling utils (takes precedence)

from dataset_headers import header_signature                      # noqa: E402
from header_reconcile import BROKEN, header_match, reconcile_cols   # noqa: E402
from refine_direct_merge_method import (                          # noqa: E402
    AS_ON_RE, DIM_ORDER, DIST_RE, MONTH_RE, NON_STATE, STATE_RE, YEAR_RE,
    header_dims, title_dims,
)
from title_components import parse_title_components                # noqa: E402

DB_PATH = str(HERE.parent / "metadata.db")
TABLE = "dublin_core_remaining"
OUT_DIR = HERE.parent
DECISIONS = HERE.parent / "collection_decisions.csv"
DECISION_FIELDS = ["new_collection_name", "new_merge_add_columns", "action",
                   "split_rule", "notes"]
SCHEMA_SUFFIX = re.compile(r" \(schema \d+\)$")

# Categorical dimensions this data needs beyond the pipeline's
# year/state/district/month/quarter vocabulary. Detected from the columns that
# separate one (schema N) sibling from another, and from the tokens that vary
# across member titles; suggested, never written, because naming the concept is
# the judgement call this script hands to a human.
EXTRA_DIMS = {
    "scheme_type": r"surface\s+(?:flow|lift)|\bflow\b.*\blift\b",
    "storage_type": r"tanks?|diversion|storage|on\s+(?:river|stream|drain|canal|tank|pond)",
    "well_type": r"dugwell|tubewell|borewell",
    "use_status": r"(?:not\s+)?in\s+use",
    "jail_type": r"\b(?:central|district|open|sub|women|special|other|borstal)\s+(?:jail|school)",
    "airline": r"air\s+\w+|airways|airlines|indigo|spicejet|jetlite|vistara",
    "social_group": r"scheduled\s+(?:caste|tribe)|\bsc\b|\bst\b",
    "urban_rural": r"\b(?:rural|urban|town|village)\b",
}

# header_dims() tokenises on camel case, so "StateName" splits but the all-caps
# "STATENAME,DISTNAME,BLOCKNAME" convention (used by the minor-irrigation census
# tables) arrives as one unsplittable token and the state/district columns go
# unseen. Augment locally rather than patching refine_direct_merge_method.py --
# that script writes merge_method off the same call, and widening it there would
# silently reclassify rows on its next run.
GLUED_GEO = {
    "state": r"^(?:state|st)(?:name|nm|cd|code)?$",
    "district": r"^(?:dist|distt|district|zila|zilla)(?:name|nm|cd|code)?$",
}

# refine_direct_merge_method.DIST_RE stops at two capitalised words, so
# three-word districts ("East Garo Hills", "South West Khasi Hills") slip past
# it. Widened here only -- that module's own matching is left as it is.
DIST_RE_WIDE = re.compile(
    r"\b([A-Z][\w.'&-]*(?:\s+[A-Z][\w.'&-]*){0,3})\s+[Dd]istrict\b")


def augment_header_dims(conforms_to: str) -> set[str]:
    """state/district columns header_dims() misses because they are glued."""
    found = set()
    for raw in (conforms_to or "").split(","):
        col = re.sub(r"[^a-z0-9]", "", raw.strip().lower())
        for dim, pat in GLUED_GEO.items():
            if re.match(pat, col):
                found.add(dim)
    return found


# Words that carry no dimension information when diffing member titles.
STOP = {
    "and", "the", "for", "from", "during", "with", "their", "its", "was", "were",
    "year", "years", "wise", "month", "months", "upto", "till", "level", "all",
    "of", "as", "on", "in", "to", "at", "by", "no", "nos", "number", "total",
    "each", "separately", "date", "end", "other", "others", "including",
}


_YEAR_TOKEN = re.compile(r"^(?:19|20)\d{2}$")


def toks(s: str) -> list[str]:
    """Comparable word tokens of a title or name."""
    return [w for w in re.findall(r"[a-z0-9]+", (s or "").lower()) if len(w) > 2]


def name_toks(s: str) -> list[str]:
    """Name tokens that must hold for every member: years are span labels, not claims."""
    return [w for w in toks(s) if not _YEAR_TOKEN.match(w)]


def varying_middles(titles: list[str]) -> list[str]:
    """
    The part of each title that is not shared by all of them.

    Trims the longest common word-prefix and word-suffix across the member
    titles; what is left is the axis the collection varies along. This is what
    the collection name must describe generically and what merge_add_columns
    must capture.
    """
    ws = [re.findall(r"\S+", t or "") for t in titles]
    ws = [w for w in ws if w]
    if len(ws) < 2:
        return []
    shortest = min(len(w) for w in ws)
    p = 0
    while p < shortest and len({w[p].lower() for w in ws}) == 1:
        p += 1
    s = 0
    while s < shortest - p and len({w[len(w) - 1 - s].lower() for w in ws}) == 1:
        s += 1
    mids = [" ".join(w[p:len(w) - s]) for w in ws]
    return sorted({m for m in mids if m})


def suggest_dims(text: str) -> list[str]:
    """Categorical dimensions EXTRA_DIMS recognises in the varying text."""
    low = text.lower()
    return sorted(d for d, pat in EXTRA_DIMS.items() if re.search(pat, low))


def pinned_geo_groups(by, sig_of):
    """
    Collections that are one series torn apart by the district each was named
    after.

    Two collections merge losslessly only if nothing distinguishes them except a
    value the data already carries. So: same header signature, a long shared
    name prefix, and the part where the names differ parses as a district/state
    that is itself a header column. "... in Vikarabad District of Telangana" and
    "... in Anantnag District of Jammu and Kashmir" qualify -- DistrictName and
    StateName are columns, so stacking them loses nothing.

    "... - Rural" vs "... - Urban", and "... for Scheduled Castes" vs "... for
    Scheduled Tribes", deliberately do NOT qualify: those differ by a value that
    exists nowhere but the title, so merging them would destroy it. They are
    correctly separate collections.
    """
    bysig = defaultdict(list)
    for coll, sig in sig_of.items():
        if sig:
            bysig[frozenset(sig)].append(coll)

    groups = {}
    for sig, colls in bysig.items():
        if len(colls) < 2:
            continue
        bare = {c: SCHEMA_SUFFIX.sub("", c) for c in colls}
        used = set()
        for i, c1 in enumerate(colls):
            if c1 in used:
                continue
            grp = [c1]
            for c2 in colls[i + 1:]:
                if c2 in used:
                    continue
                w1 = re.findall(r"\S+", bare[c1])
                w2 = re.findall(r"\S+", bare[c2])
                p = 0
                while p < min(len(w1), len(w2)) and w1[p].lower() == w2[p].lower():
                    p += 1
                if p < 5:
                    continue
                # what separates them must be geography the headers already hold
                # parse_title_components is anchored on "<DISTRICT> of <STATE>"
                # at the end of a title; these names carry a trailing "as on
                # date", so fall back to the plain "<Name> District" pattern
                # refine_direct_merge_method.py keeps for the same reason.
                g1 = dict(parse_title_components(bare[c1]))
                g2 = dict(parse_title_components(bare[c2]))
                for g, name in ((g1, bare[c1]), (g2, bare[c2])):
                    if not g.get("district"):
                        m = DIST_RE_WIDE.search(name) or DIST_RE.search(name)
                        if m:
                            g["district"] = m.group(1)
                if not (g1.get("district") and g2.get("district")):
                    continue
                if g1["district"].lower() == g2["district"].lower():
                    continue
                rest1 = " ".join(w1[p:]).lower()
                rest2 = " ".join(w2[p:]).lower()
                for g in (g1, g2):
                    for key in ("district", "state"):
                        if g.get(key):
                            rest1 = rest1.replace(g[key].lower(), " ")
                            rest2 = rest2.replace(g[key].lower(), " ")
                leftover = {w for w in re.findall(r"[a-z]+", rest1 + " " + rest2)
                            if len(w) > 2} - {"district", "state", "date",
                                              "hills", "metro", "nagar"} - STOP
                if leftover:
                    continue
                dims = header_dims(",".join(sig)) | augment_header_dims(",".join(sig))
                if "district" not in dims:
                    continue
                grp.append(c2)
                used.add(c2)
            if len(grp) > 1:
                used.add(c1)
                for c in grp:
                    groups[c] = grp
    return groups


_QUARTER_VAL = re.compile(
    r"\b(?:q[1-4]|(?:first|second|third|fourth|1st|2nd|3rd|4th)\s+quarter)\b", re.I)
_WISE = re.compile(r"\b([a-z]+)\s*-?\s*wise\b|\b(all[- ]india)\s+level\b", re.I)
_LEVEL_OF = {"state": "state", "ut": "state", "uts": "state", "district": "district",
             "block": "block", "village": "village", "town": "town", "city": "city",
             "tehsil": "sub-district", "taluka": "sub-district", "taluk": "sub-district",
             "subdistrict": "sub-district", "ward": "ward", "all india": "india",
             "all-india": "india"}


def varying_dims(members) -> set[str]:
    """
    Key dimensions whose value differs between member titles.

    Compared per member as the set of values its title names, so a span title
    ("from 2004-05 to 2013-14") repeated by every member does not count as a
    varying year.
    """
    seen = defaultdict(set)
    for m in members:
        t = m[2] or ""
        years = frozenset(re.sub(r"\s", "", y) for y in YEAR_RE.findall(t))
        seen["year"].add(years or frozenset([str(m[6] or "")]))
        seen["state"].add(frozenset(s.lower() for s in STATE_RE.findall(t)) - NON_STATE)
        dist = DIST_RE_WIDE.search(t)
        seen["district"].add(dist.group(1).lower() if dist else "")
        months = frozenset() if AS_ON_RE.search(t) else \
            frozenset(x.lower()[:3] for x in MONTH_RE.findall(t))
        seen["month"].add(months)
        seen["quarter"].add(frozenset(q.lower() for q in _QUARTER_VAL.findall(t)))
    return {d for d, vals in seen.items() if len(vals) > 1}


def admin_levels(titles) -> set[str]:
    """Distinct '<level>-wise' combinations the member titles claim."""
    out = set()
    for t in titles:
        lv = set()
        for a, b in _WISE.findall(t or ""):
            key = (a or b).lower()
            if key in _LEVEL_OF:
                lv.add(_LEVEL_OF[key])
        if lv:
            out.add("+".join(sorted(lv)))
    return out


def header_groups(members, max_cols=14):
    """
    Members grouped by reconciled column set, largest group first, as evidence.

    -> (one readable line per group, lowest match of any group against the
        largest, number of members whose 'header' is an error page)
    """
    grp = defaultdict(list)
    broken = 0
    for m in members:
        h = m[3]
        if h and BROKEN.search(h):
            broken += 1
            grp["BROKEN"].append(m)
        else:
            grp[frozenset(reconcile_cols(h)) if h else None].append(m)
    ordered = sorted(grp.items(), key=lambda kv: -len(kv[1]))
    real = [k for k, _ in ordered if isinstance(k, frozenset)]
    dom = real[0] if real else None
    lines, worst = [], None
    for i, (key, ms) in enumerate(ordered, 1):
        years = Counter(str(m[6]) for m in ms if m[6])
        yrs = ",".join(f"{y}x{c}" for y, c in sorted(years.items())) or "-"
        if isinstance(key, frozenset):
            score = 1.0 if key == dom else header_match(set(key), set(dom))
            if key != dom and score is not None:
                worst = score if worst is None else min(worst, score)
            names = sorted(reconcile_cols(next(m[3] for m in ms)).values())
            desc = (f"match={'-' if score is None else f'{score:.2f}'} "
                    f"cols({len(names)}): " + "; ".join(names[:max_cols])
                    + (" ..." if len(names) > max_cols else ""))
        else:
            desc = "NO HEADERS" if key is None else "BROKEN FILE"
        lines.append(f"G{i} [{len(ms)}] yrs={yrs} {desc} "
                     f"e.g. {min(m[2] or '' for m in ms)[:100]}")
    return lines, worst, broken


def analyse(con, table, merge_method, name_threshold, only_mac=None,
            match_threshold=0.9):
    # curate members must stack as they are, so a second header signature is a
    # defect. change_in_header members are expected to differ; what matters is
    # whether the variants reconcile with minor edits (header_reconcile.py).
    reconcile = merge_method != "curate"
    rows = con.execute(
        f'''SELECT "Collection", nid, "Title", "Conforms To", merge_add_columns,
                   "Publisher[ministry_department]", year, header_count
            FROM {table}
            WHERE merge_method = ? AND "Collection" IS NOT NULL''',
        [merge_method],
    ).fetchall()
    if only_mac is not None:
        keep = {r[0] for r in rows if (r[4] or "") == only_mac}
        rows = [r for r in rows if r[0] in keep]

    by = defaultdict(list)
    for r in rows:
        by[r[0]].append(r)

    # (schema N) sibling families: what separates one sibling from another is
    # the discriminator its name should carry instead of the bare number.
    families = defaultdict(list)
    for coll in by:
        if SCHEMA_SUFFIX.search(coll):
            families[SCHEMA_SUFFIX.sub("", coll)].append(coll)

    sig_of = {}
    for coll, members in by.items():
        first = next((m[3] for m in members if m[3]), None)
        sig_of[coll] = set(header_signature(first)) if first else set()

    geo_split = pinned_geo_groups(by, sig_of)

    unique_cols = {}
    for base, sibs in families.items():
        shared = set.intersection(*(sig_of[s] for s in sibs)) if len(sibs) > 1 else set()
        for s in sibs:
            unique_cols[s] = sorted(sig_of[s] - shared)

    out = []
    for coll, members in sorted(by.items()):
        n = len(members)
        titles = [m[2] or "" for m in members]
        headers = [m[3] for m in members]
        mac = next((m[4] for m in members if m[4]), None)
        # 'review' is a placeholder for "not decided yet", not a dimension
        mac_set = {x.strip() for x in (mac or "").split(",") if x.strip()} - {"review"}

        sigs = {header_signature(h) for h in headers if h}
        n_null = sum(1 for h in headers if not h)

        hdims = set()
        for h in headers:
            if h:
                hdims |= header_dims(h) | augment_header_dims(h)
        tdims = set()
        for t in titles:
            tdims |= title_dims(t)

        bare = SCHEMA_SUFFIX.sub("", coll)
        member_tokens = [set(toks(t)) for t in titles]
        dim_words = set()
        for d in mac_set:
            dim_words |= {w for w in re.split(r"[_\s]+", d.lower()) if len(w) > 2}
            dim_words |= {w + "s" for w in re.split(r"[_\s]+", d.lower()) if len(w) > 2}
        dim_words |= {"wise", "level", "type", "types", "class", "scope", "group",
                      "groups", "category", "categories", "basis", "measure"}
        missing = []
        for tk in sorted(set(name_toks(bare)) - dim_words):
            frac = sum(1 for mt in member_tokens if tk in mt) / n
            if frac < name_threshold:
                missing.append((tk, round(frac, 2)))
        missing.sort(key=lambda x: x[1])

        n_years = len({m[6] for m in members if m[6]})
        varies = []
        if n_years > 1:
            varies.append("year")
        if len({m[2] for m in members}) > 1:
            varies.append("title")
        # A dimension that is constant across the collection cannot be lost by
        # stacking it, so it is not a missing add-column however loudly the
        # titles name it. Only `year` is measurable here -- the table carries an
        # extracted year per row but no district column to compare.
        constant = {"year"} if n_years <= 1 else set()

        mids = varying_middles(titles)
        redundant = sorted(mac_set & hdims)
        lost = sorted((tdims - hdims) - mac_set - constant)

        groups, min_match, n_broken = header_groups(members) if reconcile else ([], None, 0)
        vary = varying_dims(members)
        suggested = [d for d in DIM_ORDER if d in vary and d not in hdims]
        suggested += [d for d in suggest_dims(" ".join(mids)) if d not in suggested]
        levels = admin_levels(titles)

        flags = []
        if missing:
            flags.append("NAME_PINS_VALUE")
        if SCHEMA_SUFFIX.search(coll):
            flags.append("SCHEMA_N_SUFFIX")
        if n_null and n_null < n:
            flags.append("HEADER_NULL_MIX")
        if reconcile:
            if min_match is not None and min_match < match_threshold:
                flags.append("HEADER_IRRECONCILABLE")
            if n_broken:
                flags.append("BROKEN_HEADER")
            if len(levels) > 1:
                flags.append("MIXED_ADMIN_LEVEL")
            if mac == "review":
                flags.append("MAC_REVIEW")
        elif len(sigs) > 1:
            flags.append("MULTI_HEADER_SIG")
        if redundant:
            flags.append("MAC_REDUNDANT")
        if lost:
            flags.append("MAC_MISSING_DIM")
        peers = [c for c in geo_split.get(coll, []) if c != coll]
        if peers:
            flags.append("SPLIT_BY_PINNED_GEO")
        if not flags:
            continue

        out.append({
            "Collection": coll,
            "defect_flags": ",".join(flags),
            "ministry": (members[0][5] or "").strip(),
            "n_members": n,
            "n_header_sigs": len(sigs),
            "n_null_headers": n_null,
            "name_tokens_missing": "; ".join(f"{t}:{f}" for t, f in missing[:8]),
            "varying_title_tokens": " | ".join(mids[:10]),
            "suggested_extra_dims": ",".join(suggest_dims(" ".join(mids))),
            "merge_with": " ;; ".join(f"[{len(by[c])}] {c}" for c in peers),
            "sibling_family": SCHEMA_SUFFIX.sub("", coll) if SCHEMA_SUFFIX.search(coll) else "",
            "n_siblings": len(families.get(SCHEMA_SUFFIX.sub("", coll), [])),
            "sibling_unique_header_cols": "; ".join(unique_cols.get(coll, [])[:8]),
            "merge_add_columns": mac or "",
            "n_distinct_year": n_years,
            "dims_that_vary": ",".join(varies),
            "title_dims": ",".join(d for d in DIM_ORDER if d in tdims),
            "header_dims": ",".join(d for d in DIM_ORDER if d in hdims),
            "mac_redundant": ",".join(redundant),
            "mac_missing": ",".join(lost),
            "header_signature": ", ".join(sorted(sigs)[0])[:400] if sigs else "",
            "sample_titles": " | ".join(sorted(set(titles))[:8]),
            "min_header_match": "" if min_match is None else round(min_match, 2),
            "admin_levels": ",".join(sorted(levels)),
            "varying_dims": ",".join(d for d in DIM_ORDER if d in vary),
            "suggested_mac": ",".join(suggested),
            "header_groups": " || ".join(groups)[:30000],
            # to fill in by hand
            "new_collection_name": "",
            "new_merge_add_columns": "",
            "action": "",
            "split_rule": "",
            "notes": "",
        })
    return out, len(by), len(rows)


COLUMNS = [
    "Collection", "defect_flags", "ministry", "n_members", "n_header_sigs",
    "n_null_headers", "min_header_match", "admin_levels",
    "name_tokens_missing", "varying_title_tokens",
    "suggested_extra_dims", "merge_with", "sibling_family", "n_siblings",
    "sibling_unique_header_cols", "merge_add_columns", "suggested_mac",
    "varying_dims", "n_distinct_year", "dims_that_vary", "title_dims",
    "header_dims", "mac_redundant", "mac_missing", "header_signature",
    "sample_titles", "header_groups", "new_collection_name",
    "new_merge_add_columns", "action", "split_rule", "notes",
]


def overlay_decisions(records, path):
    """Merge hand-written decisions onto the freshly detected rows."""
    if not path.exists():
        return 0, []
    with open(path, newline="", encoding="utf-8") as fh:
        decisions = {r["Collection"]: r for r in csv.DictReader(fh)
                     if r.get("Collection")}
    applied = 0
    for rec in records:
        d = decisions.pop(rec["Collection"], None)
        if not d:
            continue
        for f in DECISION_FIELDS:
            if d.get(f):
                rec[f] = d[f].strip()
        applied += 1
    return applied, sorted(decisions)


def write_workbook(records, xlsx_path, csv_path):
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(records)

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print(f"openpyxl unavailable - csv only: {csv_path}")
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "collection review"
    ws.append(COLUMNS)
    head = ws[1]
    fill_evidence = PatternFill("solid", fgColor="DDEBF7")
    fill_input = PatternFill("solid", fgColor="FFF2CC")
    first_input = COLUMNS.index("new_collection_name")
    for i, cell in enumerate(head):
        cell.font = Font(bold=True)
        cell.fill = fill_input if i >= first_input else fill_evidence
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for rec in records:
        ws.append([rec[c] for c in COLUMNS])
    widths = {"Collection": 60, "defect_flags": 22, "ministry": 28,
              "name_tokens_missing": 26, "varying_title_tokens": 45,
              "merge_with": 60, "sibling_family": 45,
              "sibling_unique_header_cols": 40,
              "header_signature": 45, "sample_titles": 70, "header_groups": 90,
              "new_collection_name": 45, "new_merge_add_columns": 24,
              "action": 14, "split_rule": 28, "notes": 30}
    for i, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(c, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    return xlsx_path


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=DB_PATH)
    ap.add_argument("--table", default=TABLE)
    ap.add_argument("--merge-method", default="curate")
    ap.add_argument("--name-threshold", type=float, default=0.5,
                    help="a name token present in less than this fraction of "
                         "member titles flags NAME_PINS_VALUE (default 0.5)")
    ap.add_argument("--out", default=None, help="output xlsx path")
    ap.add_argument("--decisions", default=None,
                    help="csv of hand-written decisions to overlay (default: "
                         "collection_decisions.csv for curate, "
                         "collection_decisions_<merge_method>.csv otherwise)")
    ap.add_argument("--only-mac", default=None,
                    help="only collections carrying this merge_add_columns "
                         "value, e.g. review")
    ap.add_argument("--match-threshold", type=float, default=0.9,
                    help="change_in_header: a header group matching less than "
                         "this share of the largest group's columns flags "
                         "HEADER_IRRECONCILABLE (default 0.9)")
    ap.add_argument("--summary-only", action="store_true")
    args = ap.parse_args()

    con = duckdb.connect(args.db, read_only=True)
    try:
        records, n_coll, n_rows = analyse(con, args.table, args.merge_method,
                                          args.name_threshold, args.only_mac,
                                          args.match_threshold)
    finally:
        con.close()

    flagged_rows = sum(r["n_members"] for r in records)
    print(f"table={args.table}  merge_method={args.merge_method}")
    print(f"collections            : {n_coll}  ({n_rows} rows)")
    print(f"flagged                : {len(records)}  ({flagged_rows} rows)")
    print(f"clean                  : {n_coll - len(records)}  ({n_rows - flagged_rows} rows)\n")

    per_flag = Counter()
    per_flag_rows = Counter()
    for r in records:
        for f in r["defect_flags"].split(","):
            per_flag[f] += 1
            per_flag_rows[f] += r["n_members"]
    print(f"  {'flag':<18}{'#colls':>8}{'#rows':>8}")
    for f, c in per_flag.most_common():
        print(f"  {f:<18}{c:>8}{per_flag_rows[f]:>8}")

    print(f"\n  {'ministry':<58}{'#colls':>8}")
    for m, c in Counter(r["ministry"][:56] for r in records).most_common(12):
        print(f"  {m:<58}{c:>8}")

    suffix = "" if args.merge_method == "curate" else f"_{args.merge_method}"
    decisions = Path(args.decisions) if args.decisions else \
        HERE.parent / f"collection_decisions{suffix}.csv"
    applied, orphans = overlay_decisions(records, decisions)
    decided = sum(1 for r in records if r["action"])
    print(f"\ndecisions overlaid     : {applied}  "
          f"({decided} of {len(records)} flagged rows now have an action, "
          f"{len(records) - decided} still to do)")
    if orphans:
        print(f"  ! {len(orphans)} decision(s) match no flagged collection "
              f"(already applied, or the name has changed):")
        for o in orphans[:10]:
            print(f"      {o[:100]}")

    if args.summary_only:
        return

    xlsx = Path(args.out) if args.out else OUT_DIR / f"collection_review{suffix}.xlsx"
    written = write_workbook(records, xlsx, xlsx.with_suffix(".csv"))
    print(f"\nWorkbook : {written or xlsx.with_suffix('.csv')}")
    print(f"CSV      : {xlsx.with_suffix('.csv')}")
    print("Fill new_collection_name / new_merge_add_columns / action, then run "
          "apply_collection_review.py")


if __name__ == "__main__":
    main()
