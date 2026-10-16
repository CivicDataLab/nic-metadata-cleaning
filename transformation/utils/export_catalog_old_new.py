#!/usr/bin/env python3
"""
Export the catalog metadata in the old / new layout: each catalog's raw (OGD)
value beside its mapped value, field by field.

  dublin_catalog_all_lots_old_new.xlsx  sheet "catalog_raw_vs_dublin"  <- dublin_catalog_all_lots
  dcat_catalog_all_lots_old_new.xlsx    sheet "catalog_raw_vs_dcat"    <- dcat_catalog_all_lots  (--dcat)

One header row: nid, then "old <OGD field>" / "new <column>" for every column
of the combined catalog table that data/old_new_sample_layout.csv lists for it (in
its order, under its old names; without a list, every column but merge_flag),
"old <column> (no OGD field)" where the
column was built rather than carried over (Relation from the datasets' landing
pages, new_catalog_title, the constants). old is the raw catalog's value,
catalog_metadata for lot 1 catalogs and remaining_catalog for lots 2 and 3,
matched on nid, or "NA" with no OGD field; new is the combined table's cell. A
new cell that differs from its old cell is filled yellow; a blank
prov:qualifiedAttribution is not a change.

The OGD field of each column is read off the SELECTs of
build_dublin_core_catalog_lot1.py and build_dublin_catalog_lot2n3.py, under
remaining_catalog's spelling, which is the OGD export's (catalog_metadata has
"_" for ":"). The DCAT columns are the Dublin Core ones renamed the way
export_dcat_catalog_lots_combined.py renames them.

  --nids N ...          only these catalogs
  --dataset-nids N ...  the catalogs of these datasets: the catalog whose Title is
                        the dataset's Relation[Catalog Title], in the dataset's
                        lot, the nearest nid among same-title catalogs (the rule
                        of the catalog builds)
  --unmerged-titles     new Title (dcterms:title) is the catalog's unmerged new
                        title: new_catalog_title where merge_flag is No, the
                        catalog's own Title where it was merged into a group
  --csv                 CSV instead of a workbook (UTF-8, nothing coloured, no
                        cell cut to Excel's 32,767 characters)

The new side is read from the combined tables, so re-run
export_catalog_lots_combined.py / export_dcat_catalog_lots_combined.py first if
a catalog lot changed. Nothing is written to the database. Like the other
exports it does not overwrite a workbook that is open in an office suite.

Usage:
  python3 export_catalog_old_new.py [--dcat] [--dataset-nids 604859456 603491108 ...] [--unmerged-titles]
"""

import argparse
import csv
import datetime as dt
import os
import sys
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import add_raw_sheet as raw
import export_all_lots as eal

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from build_dcat_mapping import MAPPING  # noqa: E402

q = eal.q
DB_PATH = str(Path(__file__).resolve().parent.parent / "metadata.db")
TABLES = {"dublin": ("dublin_catalog_all_lots", "catalog_raw_vs_dublin"),
          "dcat": ("dcat_catalog_all_lots", "catalog_raw_vs_dcat")}
# source_table of the combined tables -> (lot, raw table, how it spells an OGD field)
LOTS = {
    "dublin_core_catalog_lot1": ("Lot 1", "catalog_metadata", lambda f: f.replace(":", "_")),
    "dublin_catalog_lot2n3": ("Lot 2/3", "remaining_catalog", lambda f: f),
}
# Catalog column -> OGD field; "a+b" is the two fields joined.
OGD = {
    "Identifier[UUID]": "uuid",
    "Title": "title",
    "Description": "body:value",
    "Date Modified": "changed",
    "Date Created": "created",
    "Date Issued": "published_date",
    "Spatial Coverage": "field_asset_jurisdiction:name",
    "Jurisdiction": "field_ds_govt_type",
    "Creator": "field_ministry_department:name",
    "Contributor": "field_ministry_department:name",
    "Publisher": "field_ministry_department:name",
    "Subject[sector]": "field_sector:name",
    "Subject[keyword]": "keywords",
    "Relation_from_api": "from_api",
    "High Value Dataset Category": "high_value_dataset",
    "unmappable_is_api_available": "is_api_available",
    "unmappable_ogdp_view_count": "ogdp_view_count",
    "Identifier": "ogpl_module_domain_name+node_alias",
}
# The renames of export_dcat_catalog_lots_combined.py (that script exports on import).
DCAT_EXTRA = {"Contributor": "dcterms:contributor", "Identifier": "dcat:landingPage"}
DCAT_KEEP = {"Subject[sector]"}
NO_OGD, NA = " (no OGD field)", "NA"
# Not shown: the user deleted the lot, dataset nid and merge_flag columns from the
# sheet (2026-10-14).
LEAVE_OUT = {"nid", "source_table", "merge_flag"}
# The columns, their order and the "old" names of the user's own edit of the
# 6-catalog sheets (2026-10-15), by sheet; shared with export_all_lots_old_new.py.
LAYOUT = Path(eal.REPO) / "data" / "old_new_sample_layout.csv"
CELL_MAX = 32767
BOLD = Font(bold=True)
CENTRE = Alignment(horizontal="center")
FILLS = {k: PatternFill("solid", fgColor=f"FF{rgb}")
         for k, rgb in (("old", "D9D9D9"), ("new", "C6E0B4"), ("changed", "FFF2CC"))}


def dublin_names(con, table: str) -> dict:
    """Column of the combined table -> the Dublin Core catalog column behind it."""
    cols = [c for c, *_ in con.execute(f"DESCRIBE {q(table)}").fetchall()]
    if table == TABLES["dublin"][0]:
        return {c: c for c in cols}
    dublin = [c for c, *_ in con.execute(f"DESCRIBE {q(TABLES['dublin'][0])}").fetchall()]
    renamed = {DCAT_EXTRA.get(c) or ("" if c in DCAT_KEEP else MAPPING.get(c, ("",))[0]) or c: c
               for c in dublin if c != "Publisher"}
    out = {}
    for c in cols:
        if c == "dcterms:publisher" or c.startswith("prov:qualifiedAttribution["):
            out[c] = "Publisher"
        elif c == "dcat:hadRole":          # a constant, no Dublin Core column or OGD field
            out[c] = c
        elif c in renamed:
            out[c] = renamed[c]
        else:
            sys.exit(f"{table}.{c}: no Dublin Core catalog column behind it")
    return out


CUT = {}          # column -> cells cut to Excel's cell limit


def text(v, column: str, excel: bool = True) -> str:
    v = "" if v is None else str(v)
    if not excel:
        return v
    v = ILLEGAL_CHARACTERS_RE.sub("", v)
    if len(v) <= CELL_MAX:
        return v
    CUT[column] = CUT.get(column, 0) + 1
    return v[:CELL_MAX - 3] + "..."


def pick_catalogs(con, args) -> list[int]:
    """-> catalog nids, or [] for all."""
    if args.nids:
        return [int(n) for n in args.nids]
    if not args.dataset_nids:
        return []
    nids = ", ".join(repr(n) for n in args.dataset_nids)
    rows = con.execute(f"""
        WITH d AS (SELECT lot, nid, trim("Relation[Catalog Title]") AS ct
                   FROM dublin_core_all_lots WHERE nid IN ({nids})),
             c AS (SELECT nid AS cat_nid, trim("Title") AS t, source_table FROM dublin_catalog_all_lots)
        SELECT d.nid, c.cat_nid FROM d JOIN c ON c.t = d.ct
               AND (c.source_table = 'dublin_core_catalog_lot1') = (d.lot = 'Lot 1')
        QUALIFY row_number() OVER (PARTITION BY d.nid
                                   ORDER BY abs(c.cat_nid - CAST(d.nid AS BIGINT)), c.cat_nid) = 1
        """).fetchall()
    found = dict(rows)
    if missing := [n for n in args.dataset_nids if n not in found]:
        print(f"! no catalog found for datasets {missing}")
    if not found:
        sys.exit("none of the datasets has a catalog")
    return list(dict.fromkeys(found[n] for n in args.dataset_nids if n in found))


def save_xlsx(path, sheet_name: str, head: list, body: list):
    """The sheet: bold old / new header, frozen nid column, changed new cells yellow."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "B2"
    last = get_column_letter(len(head))
    ws.auto_filter.ref = f"A1:{last}{len(body) + 1}"
    for i, h in enumerate(head, 1):
        longest = max([len(h)] + [len(line[i - 1][0] if isinstance(line[i - 1], tuple) else line[i - 1])
                                  for line in body])
        ws.column_dimensions[get_column_letter(i)].width = 12 if i <= 1 else min(max(longest + 2, 10), 50)

    def cell(v, fill=None, bold=False):
        c = WriteOnlyCell(ws, value=v)
        if fill:
            c.fill = FILLS[fill]
        if bold:
            c.font, c.alignment = BOLD, CENTRE
        return c

    ws.append([cell(h, h.split(" ", 1)[0] if h[:4] in ("old ", "new ") else None, bold=True) for h in head])
    for line in body:
        ws.append([cell(v[0], "changed" if v[1] else None) if isinstance(v, tuple) else cell(v) for v in line])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=DB_PATH)
    ap.add_argument("--dcat", action="store_true", help="raw vs DCAT, from dcat_catalog_all_lots")
    group = ap.add_mutually_exclusive_group()
    group.add_argument("--nids", nargs="+", metavar="NID", help="only these catalogs")
    group.add_argument("--dataset-nids", nargs="+", metavar="NID", help="the catalogs of these datasets")
    ap.add_argument("--unmerged-titles", action="store_true",
                    help="new Title (dcterms:title) = new_catalog_title where merge_flag is No")
    ap.add_argument("--csv", action="store_true",
                    help="write a CSV (no yellow, no cell cut to Excel's limit) instead of a workbook")
    ap.add_argument("--out", help="output path (default: <table name>_old_new[_<n>_nids].xlsx / .csv in the repo root)")
    ap.add_argument("--force", action="store_true",
                    help="overwrite even if the file looks open in an office suite")
    args = ap.parse_args()
    given = args.nids or args.dataset_nids or []
    if not all(n.isdigit() for n in given):
        sys.exit(f"nids are numbers: {given}")

    table, sheet_name = TABLES["dcat" if args.dcat else "dublin"]
    ext = ".csv" if args.csv else ".xlsx"
    default = f"{table}_old_new_{len(given)}_nids{ext}" if given else f"{table}_old_new{ext}"
    out = Path(args.out) if args.out else eal.REPO / default
    holder = eal.wrapper.file_is_open(out)
    if holder and not args.force:
        out = out.with_name(f"{out.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{out.suffix}")
        print(f"! the workbook is open in {holder} - writing {out.name} instead (--force to override).")

    con = duckdb.connect(args.db, read_only=True)
    try:
        cat_nids = pick_catalogs(con, args)
        behind = dublin_names(con, table)
        fields = [c for c in behind if c not in LEAVE_OUT]
        with open(LAYOUT, encoding="utf-8") as fh:
            layout = {r["new_column"]: r["old_label"] for r in csv.DictReader(fh) if r["sheet"] == table}
        if layout:
            if missing := [f for f in layout if f not in fields]:
                print(f"! in {LAYOUT.name}, not in {table}: {missing}")
            print(f"  left out, as in {LAYOUT.name}: {[f for f in fields if f not in layout]}")
            fields = [f for f in layout if f in fields]
        new_expr = {f: f"m.{q(f)}" for f in fields}
        if args.unmerged_titles:
            for f in fields:
                if behind[f] == "Title":
                    new_expr[f] = f"CASE WHEN m.merge_flag = 'No' THEN m.new_catalog_title ELSE m.{q(f)} END"
        only = f" AND m.nid IN ({', '.join(map(str, cat_nids))})" if cat_nids else ""

        labels, rows = {}, []
        for source, (lot, raw_table, spell) in LOTS.items():
            described = [(n, t) for n, t, *_ in con.execute(f"DESCRIBE {q(raw_table)}").fetchall()]
            have = {n for n, _ in described}
            raw_select = ("SELECT " + ", ".join(f"{raw.text_expr(con, raw_table, n, t)} AS {q(n)}"
                                               for n, t in described) + f" FROM {q(raw_table)}")
            old = {}
            for f in fields:
                ogd = OGD.get(behind[f])
                parts = [spell(x) for x in ogd.split("+")] if ogd else []
                if parts and set(parts) <= have:
                    old[f] = " || ".join(f"coalesce(r.{q(x)}, '')" for x in parts) if len(parts) > 1 \
                        else f"r.{q(parts[0])}"
                    labels.setdefault(f, {})[lot] = ogd
            select = ["m.nid"]
            for f in fields:
                select += [old.get(f, "NULL"), f"CAST({new_expr[f]} AS VARCHAR)"]
            got = con.execute(f"""
                SELECT {', '.join(select)} FROM {q(table)} m
                LEFT JOIN ({raw_select}) r ON CAST(r.nid AS BIGINT) = m.nid
                WHERE m.source_table = '{source}'{only} ORDER BY m.rowid""").fetchall()
            rows += [(lot, set(old), r) for r in got]
            unmatched = con.execute(f"""SELECT count(*) FROM {q(table)} m WHERE m.source_table = '{source}'{only}
                AND NOT EXISTS (SELECT 1 FROM {q(raw_table)} r WHERE CAST(r.nid AS BIGINT) = m.nid)""").fetchone()[0]
            print(f"  {lot}: {len(got)} catalogs, old from {raw_table} ({len(old)} of {len(fields)} "
                  f"columns have an OGD field), {unmatched} with no raw row")
    finally:
        con.close()
    if cat_nids and len(rows) != len(cat_nids):
        sys.exit(f"{len(rows)} of the {len(cat_nids)} catalogs found in {table}")

    present = {lot for lot, _, _ in rows}
    head = ["nid"]
    for f in fields:
        ogd = " | ".join(dict.fromkeys(v for lot, v in labels.get(f, {}).items() if lot in present))
        head += [f"old {layout.get(f) or ogd or f + NO_OGD}", f"new {f}"]

    body, changed, na = [], 0, 0
    for lot, has_old, r in rows:
        line = [str(r[0])]
        for i, f in enumerate(fields):
            old = text(r[1 + 2 * i], f"old {f}", not args.csv)
            new = text(r[2 + 2 * i], f"new {f}", not args.csv)
            if f not in has_old:
                na += 1
            quiet = f.startswith("prov:qualifiedAttribution") and not new.strip()
            yellow = old.strip() != new.strip() and not quiet
            changed += yellow
            line += [old if f in has_old else NA, (new, yellow)]
        body.append(line)

    tmp = out.with_name(f".{out.stem}.{os.getpid()}.tmp")
    try:
        if args.csv:
            with open(tmp, "w", encoding="utf-8", newline="") as fh:
                w = csv.writer(fh, lineterminator="\n")
                w.writerow(head)
                w.writerows([v[0] if isinstance(v, tuple) else v for v in line] for line in body)
        else:
            save_xlsx(tmp, sheet_name, head, body)
        os.replace(tmp, out)
    finally:
        tmp.unlink(missing_ok=True)
    print(f"{'CSV' if args.csv else f'Sheet {sheet_name!r}'}: {len(body)} catalogs x {len(head)} columns; "
          f"changed (yellow) cells: {changed}; NA cells: {na}")
    if CUT:
        print(f"  cut to {CELL_MAX} characters, ending in '...': {CUT}")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
