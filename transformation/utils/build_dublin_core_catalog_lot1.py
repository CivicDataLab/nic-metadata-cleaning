"""
Build dublin_core_catalog_lot1: the catalog-level Dublin Core table for lot 1.

Source is the raw catalog_metadata table (371 data.gov.in catalogs, one row per
nid). Columns follow data/dublin_core_transformed/dublin_core_catalog_mapping.csv,
with these adjustments:

  - nid and Identifier[UUID] are kept as keys.
  - new_catalog_title is the merged catalog name already set on the lot 1
    datasets (dublin_core_metadata.new_catalog_title, joined on the trimmed
    Relation[Catalog Title]). Every catalog maps to exactly one name; the
    script aborts if that stops being true.
  - The CSV's single "Subject" glued sector and keywords together with no
    separator, so it is split into Subject[sector] / Subject[keyword] (the
    dataset tables' naming).
  - Subject[keyword] is not the catalog's own keywords (those stay in
    catalog_metadata). It is chosen from the dataset keywords
    (dublin_core_metadata.Subject[keyword], comma-separated) of all datasets
    under the same new_catalog_title, so every catalog under one title gets
    the same list:
      * synonyms in KEYWORD_GROUPS collapse to the group's first term;
      * keywords are ranked by how many original catalogs under the title
        use them (then by dataset count), so one very large catalog cannot
        dominate;
      * a keyword used by GENERIC_MIN_TITLES or more of the new titles is
        generic and goes to the bottom, unless at least half of this title's
        original catalogs use it (keeps "survey" for the survey title);
      * only the top KEYWORD_SLOTS keywords are kept.
  - Spatial Coverage comes from the Title, not the raw
    field_asset_jurisdiction_name (which says "All India;State" even for
    single-state HMIS catalogs): the State/UT named in the title, or "India"
    when the title names none or speaks of several States ("EAG States").
  - Dates are ISO YYYY-MM-DD and Identifier is a full https URL, matching
    dublin_core_metadata.
  - Relation columns follow catalog_metadata_mapping_table in
    data/sample_catalog_metadata/NIC_metadata_curation_master.xlsx: Relation
    links the catalog to its resources, the qualified Relation[...] columns
    link catalogs to catalogs.
      * Relation: the data.gov.in page of every dataset in the catalog,
        "; "-separated in dataset nid order. Built from
        raw_metadata.node_alias; ~4.2k aliases are stored as
        "resources/<slug>", which data.gov.in answers with Page Not Found, so
        they are rewritten to /resource/;
      * Relation[Has Part]: NULL, no lot 1 catalog contains another catalog;
      * Relation[Is Part Of]: the merged catalog, new_catalog_title;
      * Relation[Replaces] / Relation[Is Replaced By]: six titles exist as a
        2013 and a 2014 catalog, the 2014 one republishing the same data. The
        newer replaces the older;
      * Relation[Is Referenced By] / [Is Version Of] / [Has Version]: no source
        for them yet, left NULL.
    Datasets point at their catalog by title only. For the duplicate titles
    each dataset goes to the catalog with the nearest nid: on data.gov.in a
    dataset's nid sits right below its catalog's (86931 -> 86932).

Creates a new table only; nothing existing is modified. Backs up the DB to
metadata.db.pre_catalog_lot1_*.bak before writing (repo convention).

Usage:
  python build_dublin_core_catalog_lot1.py --dry-run            # preview, no write
  python build_dublin_core_catalog_lot1.py                      # create the table
  python build_dublin_core_catalog_lot1.py --replace            # drop and rebuild
  python build_dublin_core_catalog_lot1.py --replace --export   # ...and write the xlsx
  python build_dublin_core_catalog_lot1.py --export-only        # xlsx from the current table
"""

import argparse
import re
import shutil
import time
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DB_PATH = str(HERE.parent / "metadata.db")
SOURCE = "catalog_metadata"
DATASETS = "dublin_core_metadata"
RAW_DATASETS = "raw_metadata"
TABLE = "dublin_core_catalog_lot1"
XLSX_PATH = HERE.parent.parent / "data" / "dublin_core_transformed" / f"{TABLE}.xlsx"
EXCEL_CELL_MAX = 32767

NEW_TITLES_SQL = f"""
    SELECT trim("Relation[Catalog Title]") AS catalog_title,
           count(DISTINCT new_catalog_title) AS n_titles,
           any_value(new_catalog_title)      AS new_catalog_title
    FROM {DATASETS}
    GROUP BY 1
"""

KEYWORD_SLOTS = 15
GENERIC_MIN_TITLES = 7

# Checked by hand against dublin_core_metadata: each group is tagged on the
# same (or nearly the same) datasets and means the same thing. The first term
# is the one written to Subject[keyword]; the rest are dropped as synonyms.
# Pairs that are always co-tagged but mean different things stay separate
# (birth rate / death rate, alcohol use / tobacco use, health indicators /
# performance).
KEYWORD_GROUPS = [
    ["ranking", "comparison"],
    ["inpatient", "hospital admission"],
    ["outpatient", "hospital visit"],
    ["ayush", "traditional medicine"],
    ["statistics", "sample data"],
    ["health worker", "medical staff", "staffing"],
    ["immunisation", "vaccination"],
    ["health infrastructure", "facilities"],
    ["sub centre", "health post"],
    ["government hospital", "hospital"],
    ["health expenditure", "budget"],
    ["medical education", "medical college"],
    ["asthma", "respiratory illness"],
    ["goitre", "thyroid disorder"],
    ["elderly", "ageing"],
    ["health mission", "health scheme"],
    ["child health", "children", "kids health"],
    ["family planning", "contraceptive", "birth control"],
    ["rural", "village health"],
    ["scheduled tribe", "tribal areas"],
    ["health data", "health report"],
    ["women", "female health"],
]

NEW_KEYWORDS_SQL = f"""
    WITH kw AS (
        SELECT DISTINCT new_catalog_title, trim("Relation[Catalog Title]") AS cat,
               nid, trim(k) AS keyword
        FROM {DATASETS}, unnest(string_split("Subject[keyword]", ',')) AS s(k)
        WHERE trim(k) <> ''
    ), kg AS (
        SELECT kw.*, coalesce(g.grp, kw.keyword) AS grp
        FROM kw LEFT JOIN keyword_groups g USING (keyword)
    ), title_cats AS (
        SELECT new_catalog_title, count(DISTINCT cat) AS n_cats FROM kw GROUP BY 1
    ), spread AS (
        SELECT grp, count(DISTINCT new_catalog_title) AS n_titles FROM kg GROUP BY 1
    ), slots AS (
        SELECT new_catalog_title, grp, count(DISTINCT cat) AS cats, count(DISTINCT nid) AS ds
        FROM kg GROUP BY 1, 2
    ), ranked AS (
        SELECT s.new_catalog_title, s.grp,
               row_number() OVER (
                   PARTITION BY s.new_catalog_title
                   ORDER BY (sp.n_titles >= {GENERIC_MIN_TITLES} AND s.cats * 2 < tc.n_cats),
                            s.cats DESC, s.ds DESC, s.grp
               ) AS slot
        FROM slots s JOIN spread sp USING (grp) JOIN title_cats tc USING (new_catalog_title)
    ), raw AS (
        SELECT new_catalog_title, count(DISTINCT keyword) AS n_raw FROM kw GROUP BY 1
    )
    SELECT r.new_catalog_title,
           any_value(n_raw)                                  AS n_raw,
           count(*)                                          AS n_keywords,
           count(*) FILTER (WHERE r.slot <= {KEYWORD_SLOTS}) AS n_kept,
           string_agg(r.grp, ', ' ORDER BY r.slot)
               FILTER (WHERE r.slot <= {KEYWORD_SLOTS})      AS keywords
    FROM ranked r JOIN raw USING (new_catalog_title)
    GROUP BY 1
"""

# States/UTs as the lot 1 catalog titles name them. Dadra and Nagar Haveli and
# Daman and Diu stay separate: the catalogs predate their 2020 merger.
STATES = [
    "Andaman and Nicobar Islands", "Andhra Pradesh", "Arunachal Pradesh", "Assam",
    "Bihar", "Chandigarh", "Chhattisgarh", "Dadra and Nagar Haveli", "Daman and Diu",
    "Delhi", "Goa", "Gujarat", "Haryana", "Himachal Pradesh", "Jammu and Kashmir",
    "Jharkhand", "Karnataka", "Kerala", "Ladakh", "Lakshadweep", "Madhya Pradesh",
    "Maharashtra", "Manipur", "Meghalaya", "Mizoram", "Nagaland", "Odisha",
    "Puducherry", "Punjab", "Rajasthan", "Sikkim", "Tamil Nadu", "Telangana",
    "Tripura", "Uttar Pradesh", "Uttarakhand", "West Bengal",
]
# Other spellings, in normalised form ("&" is already "and"; see _norm).
STATE_ALIASES = {
    "a and n islands": "Andaman and Nicobar Islands",
    "andaman and nicobar island": "Andaman and Nicobar Islands",
    "nct of delhi": "Delhi",
    "orissa": "Odisha",
    "pondicherry": "Puducherry",
    "uttaranchal": "Uttarakhand",
}


def _norm(text: str) -> str:
    """Lower-case, "&" -> "and", punctuation -> space, whitespace collapsed."""
    s = re.sub(r"[^a-z0-9]+", " ", text.lower().replace("&", " and "))
    return re.sub(r"\s+", " ", s).strip()


_STATE_LOOKUP = {_norm(s): s for s in STATES} | STATE_ALIASES
# Longest first, so "Andaman and Nicobar Islands" wins over a shorter alias.
_STATE_RE = re.compile(
    r"\b(" + "|".join(re.escape(k) for k in sorted(_STATE_LOOKUP, key=len, reverse=True)) + r")\b"
)


# A title that names a State but also a group of States covers several of
# them, e.g. "Top 10 causes of death in EAG States- Assam and other state".
_MULTI_STATE_RE = re.compile(r"\b(states|other state)\b")


def spatial_from_title(title: str) -> str:
    """The State/UT(s) named in a catalog title, else "India"."""
    norm = _norm(title or "")
    if _MULTI_STATE_RE.search(norm):
        return "India"
    found = []
    for m in _STATE_RE.finditer(norm):
        state = _STATE_LOOKUP[m.group(1)]
        if state not in found:
            found.append(state)
    return "; ".join(found) or "India"


# One row per lot 1 dataset with its catalog nid and page link.
DATASET_LINKS_SQL = f"""
    WITH ds AS (
        SELECT d.nid::BIGINT                     AS dataset_nid,
               trim(d."Relation[Catalog Title]") AS catalog_title,
               CASE
                   WHEN r.node_alias LIKE '/resource/%'
                       THEN 'https://data.gov.in' || r.node_alias
                   WHEN r.node_alias LIKE 'resources/%'
                       THEN 'https://data.gov.in/resource/' || substr(r.node_alias, 11)
               END                               AS dataset_link,
               r.node_alias LIKE 'resources/%'   AS link_fixed
        FROM {DATASETS} d JOIN {RAW_DATASETS} r ON r.nid::VARCHAR = d.nid::VARCHAR
    ), matched AS (
        SELECT ds.*, c.nid AS catalog_nid,
               row_number() OVER (
                   PARTITION BY ds.dataset_nid ORDER BY abs(c.nid - ds.dataset_nid), c.nid
               ) AS rn
        FROM ds JOIN {SOURCE} c ON trim(c.title) = ds.catalog_title
    )
    SELECT * EXCLUDE (rn) FROM matched WHERE rn = 1
"""

CATALOG_RELATIONS_SQL = """
    SELECT catalog_nid,
           string_agg(dataset_link, '; ' ORDER BY dataset_nid) AS relation
    FROM dataset_links
    GROUP BY 1
"""

# Same-title catalogs in publication order; each edition replaces the one before.
CATALOG_EDITIONS_SQL = f"""
    WITH e AS (
        SELECT nid, trim(title) AS title,
               'https://' || ogpl_module_domain_name || node_alias AS link,
               row_number() OVER (PARTITION BY trim(title) ORDER BY published_date, nid) AS ed
        FROM {SOURCE}
    )
    SELECT e.nid, prev.link AS replaces, nxt.link AS is_replaced_by
    FROM e
    LEFT JOIN e prev ON prev.title = e.title AND prev.ed = e.ed - 1
    LEFT JOIN e nxt  ON nxt.title  = e.title AND nxt.ed  = e.ed + 1
"""

SELECT_SQL = f"""
    SELECT
        c.nid                                                   AS "nid",
        c.uuid                                                  AS "Identifier[UUID]",
        trim(c.title)                                           AS "Title",
        t.new_catalog_title                                     AS "new_catalog_title",
        c.body_value                                            AS "Description",
        strftime(c.changed, '%Y-%m-%d')                         AS "Date Modified",
        strftime(c.created, '%Y-%m-%d')                         AS "Date Created",
        strftime(c.published_date, '%Y-%m-%d')                  AS "Date Issued",
        title_spatial(trim(c.title))                            AS "Spatial Coverage",
        c.field_ds_govt_type                                    AS "Jurisdiction",
        c.field_ministry_department_name                        AS "Creator",
        c.field_ministry_department_name                        AS "Contributor",
        c.field_ministry_department_name                        AS "Publisher",
        c.field_sector_name                                     AS "Subject[sector]",
        k.keywords                                              AS "Subject[keyword]",
        c.from_api                                              AS "Relation_from_api",
        c.high_value_dataset                                    AS "High Value Dataset Category",
        c.is_api_available                                      AS "unmappable_is_api_available",
        c.ogdp_view_count                                       AS "unmappable_ogdp_view_count",
        'https://' || c.ogpl_module_domain_name || c.node_alias AS "Identifier",
        rel.relation                                            AS "Relation",
        'NDSAP'                                                 AS "Rights Statement",
        'Open'                                                  AS "Access Rights",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Has Part]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Is Referenced By]",
        e.is_replaced_by                                        AS "Relation[Is Replaced By]",
        e.replaces                                              AS "Relation[Replaces]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Is Version Of]",
        CAST(NULL AS VARCHAR)                                   AS "Relation[Has Version]",
        t.new_catalog_title                                     AS "Relation[Is Part Of]"
    FROM {SOURCE} c
    LEFT JOIN new_titles t ON t.catalog_title = trim(c.title)
    LEFT JOIN new_keywords k ON k.new_catalog_title = t.new_catalog_title
    LEFT JOIN catalog_relations rel ON rel.catalog_nid = c.nid
    LEFT JOIN catalog_editions e ON e.nid = c.nid
    ORDER BY c.nid
"""


def check(con: duckdb.DuckDBPyConnection) -> None:
    """Set up the temp objects SELECT_SQL needs and print what they resolve to.

    Aborts unless every catalog maps to exactly one new_catalog_title.
    """
    con.create_function("title_spatial", spatial_from_title, ["VARCHAR"], "VARCHAR")
    con.execute(f"CREATE OR REPLACE TEMP VIEW new_titles AS {NEW_TITLES_SQL}")
    total, missing, ambiguous = con.execute(f"""
        SELECT count(*),
               count(*) FILTER (WHERE t.catalog_title IS NULL OR t.new_catalog_title IS NULL),
               count(*) FILTER (WHERE t.n_titles > 1)
        FROM {SOURCE} c LEFT JOIN new_titles t ON t.catalog_title = trim(c.title)
    """).fetchone()
    print(f"{SOURCE}: {total} catalogs; no new_catalog_title: {missing}; "
          f"more than one: {ambiguous}")
    if missing or ambiguous:
        raise SystemExit("new_catalog_title mapping is not one-to-one; aborting.")
    con.execute("CREATE OR REPLACE TEMP TABLE keyword_groups (keyword VARCHAR, grp VARCHAR)")
    con.executemany(
        "INSERT INTO keyword_groups VALUES (?, ?)",
        [(k, g[0]) for g in KEYWORD_GROUPS for k in g],
    )
    con.execute(f"CREATE OR REPLACE TEMP VIEW new_keywords AS {NEW_KEYWORDS_SQL}")
    print("\nKeywords per new_catalog_title (kept / after synonyms / raw unique):")
    for title, raw, n, kept in con.execute(
        "SELECT new_catalog_title, n_raw, n_keywords, n_kept FROM new_keywords "
        "ORDER BY n_raw DESC, new_catalog_title"
    ).fetchall():
        print(f"  {kept:3} / {n:3} / {raw:3}  {title}")
    india, one_state, several = con.execute(f"""
        SELECT count(*) FILTER (WHERE s = 'India'),
               count(*) FILTER (WHERE s <> 'India' AND s NOT LIKE '%;%'),
               count(*) FILTER (WHERE s LIKE '%;%')
        FROM (SELECT title_spatial(trim(title)) AS s FROM {SOURCE})
    """).fetchone()
    print(f"\nSpatial Coverage from titles: India {india}; one State/UT {one_state}; "
          f"several {several}")

    con.execute(f"CREATE OR REPLACE TEMP TABLE dataset_links AS {DATASET_LINKS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP VIEW catalog_relations AS {CATALOG_RELATIONS_SQL}")
    con.execute(f"CREATE OR REPLACE TEMP VIEW catalog_editions AS {CATALOG_EDITIONS_SQL}")
    n_datasets = con.execute(f"SELECT count(*) FROM {DATASETS}").fetchone()[0]
    linked, fixed, no_link, catalogs, split = con.execute(f"""
        SELECT count(*), count(*) FILTER (WHERE link_fixed),
               count(*) FILTER (WHERE dataset_link IS NULL),
               count(DISTINCT catalog_nid),
               count(*) FILTER (WHERE catalog_title IN (
                   SELECT trim(title) FROM {SOURCE} GROUP BY 1 HAVING count(*) > 1))
        FROM dataset_links
    """).fetchone()
    replaced = con.execute(
        "SELECT count(*) FROM catalog_editions WHERE is_replaced_by IS NOT NULL"
    ).fetchone()[0]
    print(f"\nDatasets matched to a catalog: {linked} of {n_datasets}, across {catalogs} "
          f"catalogs ({split} split between same-title catalogs by nearest nid)")
    print(f"Relation (dataset) links: {linked - no_link} ({fixed} rewritten from "
          f"'resources/'); datasets without a page alias: {no_link}")
    print(f"Catalogs replaced by a newer same-title catalog: {replaced}")
    if linked != n_datasets:
        raise SystemExit("Some datasets did not match a catalog; aborting.")


def export_xlsx() -> None:
    """Write TABLE to XLSX_PATH: the table on one sheet, dataset links on another.

    Excel cells hold at most EXCEL_CELL_MAX characters and Relation runs past
    1 MB for the big HMIS catalogs, so the "Relation" sheet lists every
    catalog -> dataset link and oversize cells on the main sheet point there
    instead.
    """
    con = duckdb.connect(DB_PATH, read_only=True)
    try:
        df = con.execute(f'SELECT * FROM {TABLE} ORDER BY new_catalog_title, "Title", nid').df()
    finally:
        con.close()

    links = df[["nid", "Title", "Relation"]].dropna(subset=["Relation"])
    links = (links.assign(Relation=links["Relation"].str.split("; "))
                  .explode("Relation")
                  .rename(columns={"nid": "catalog_nid", "Title": "catalog_title",
                                   "Relation": "dataset_link"}))

    relation = df["Relation"].fillna("")
    too_long = relation.str.len() > EXCEL_CELL_MAX
    counts = relation.str.count("; ") + 1
    df.loc[too_long, "Relation"] = (
        counts[too_long].astype(str) + " datasets; full list on sheet 'Relation'"
    )
    longest = df.astype(str).apply(lambda s: s.str.len().max())
    if longest.max() > EXCEL_CELL_MAX:
        raise SystemExit(f"{longest.idxmax()} still exceeds the Excel cell limit.")

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as xw:
        for sheet, frame in ((TABLE, df), ("Relation", links)):
            frame.to_excel(xw, sheet_name=sheet, index=False)
            ws = xw.sheets[sheet]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            widths = frame.astype(str).apply(lambda s: s.str.len().max())
            for i, col in enumerate(frame.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = (
                    min(60, max(len(col), int(widths[col])) + 2))
    print(f"Wrote {XLSX_PATH}: {len(df)} catalogs ({int(too_long.sum())} Relation cells "
          f"moved to sheet 'Relation'), {len(links)} dataset links.")


def main(dry_run: bool, replace: bool, export: bool, export_only: bool) -> None:
    if export_only:
        export_xlsx()
        return
    con = duckdb.connect(DB_PATH, read_only=True)
    check(con)
    exists = con.execute(
        "SELECT count(*) FROM information_schema.tables WHERE table_name = ?", [TABLE]
    ).fetchone()[0]
    if dry_run:
        print("\n[dry-run] No changes written.")
        con.close()
        return
    con.close()
    if exists and not replace:
        raise SystemExit(f"{TABLE} already exists; pass --replace to rebuild it.")

    dst = f"{DB_PATH}.pre_catalog_lot1_{time.strftime('%Y%m%d_%H%M%S')}.bak"
    shutil.copy2(DB_PATH, dst)
    print(f"Backed up DB -> {dst}")

    con = duckdb.connect(DB_PATH)
    try:
        check(con)
        verb = "CREATE OR REPLACE TABLE" if replace else "CREATE TABLE"
        con.execute(f'{verb} {TABLE} AS {SELECT_SQL}')
        n = con.execute(f"SELECT count(*) FROM {TABLE}").fetchone()[0]
        print(f"Created {TABLE}: {n} rows.")
    finally:
        con.close()
    if export:
        export_xlsx()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=f"Build {TABLE} from {SOURCE}.")
    parser.add_argument("--dry-run", action="store_true", help="Preview without writing.")
    parser.add_argument("--replace", action="store_true", help=f"Drop and rebuild {TABLE}.")
    parser.add_argument("--export", action="store_true", help=f"Write {XLSX_PATH.name} after building.")
    parser.add_argument("--export-only", action="store_true",
                        help=f"Write {XLSX_PATH.name} from the current table; no build.")
    args = parser.parse_args()
    main(args.dry_run, args.replace, args.export, args.export_only)
