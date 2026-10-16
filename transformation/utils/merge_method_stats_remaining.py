"""merge_method roll-up for dublin_core_remaining.xlsx -> data/merge_method_stats_remaining.csv

One row per merge_method with the old-vs-new dataset and catalog counts:

old datasets : one row == one old dataset (nid is unique).
new datasets : mergeable rows collapse into their `Collection`; rows without a
               Collection stay standalone datasets; 'to be removed' rows are
               retired, so they contribute 0.
old catalogs : distinct "Relation[Catalog Title]" (the pre-merge catalogs).
new catalogs : distinct `new_catalog_title`.

Both catalog columns are DISTINCT counts, so the TOTAL row is its own distinct
count over every row -- summing the per-method values double-counts catalogs
that span merge methods.
"""

import csv
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "..", "dublin_core_remaining.xlsx")
OUT = os.path.join(HERE, "..", "..", "data", "merge_method_stats_remaining.csv")

DROPPED_METHODS = {"to be removed"}
# Print order; anything unexpected is appended after these, before TOTAL.
METHOD_ORDER = ["non-mergeable", "curate", "direct", "change_in_header",
                "to be removed", "(not classified)"]

COLS = ("merge_method", "Collection", "new_catalog_title",
        "Relation[Catalog Title]")


def clean(v):
    """Excel nulls reach us as None or as the literal strings ''/'None'."""
    if v is None:
        return None
    v = str(v).strip()
    return None if v in ("", "None") else v


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {c: header.index(c) for c in COLS}

    stats = {}  # merge_method -> counters
    all_old_cats, all_new_cats = set(), set()

    for row in rows:
        mm = clean(row[idx["merge_method"]]) or "(not classified)"
        s = stats.setdefault(mm, {"old": 0, "colls": set(), "standalone": 0,
                                  "old_cats": set(), "new_cats": set()})
        s["old"] += 1
        coll = clean(row[idx["Collection"]])
        if coll is not None:
            s["colls"].add(coll)
        elif mm not in DROPPED_METHODS:
            s["standalone"] += 1
        oc = clean(row[idx["Relation[Catalog Title]"]])
        if oc is not None:
            s["old_cats"].add(oc)
            all_old_cats.add(oc)
        nc = clean(row[idx["new_catalog_title"]])
        if nc is not None:
            s["new_cats"].add(nc)
            all_new_cats.add(nc)

    ordered = [m for m in METHOD_ORDER if m in stats]
    ordered += [m for m in stats if m not in METHOD_ORDER]

    out = []
    for mm in ordered:
        s = stats[mm]
        new = 0 if mm in DROPPED_METHODS else len(s["colls"]) + s["standalone"]
        out.append([mm, s["old"], new, len(s["old_cats"]), len(s["new_cats"])])

    out.append(["TOTAL",
                sum(r[1] for r in out),
                sum(r[2] for r in out),
                len(all_old_cats),
                len(all_new_cats)])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["merge_method", "old datasets", "new datasets",
                    "old catalogs", "new catalogs"])
        w.writerows(out)

    width = max(len(r[0]) for r in out)
    for r in out:
        print(f"{r[0]:<{width}}  {r[1]:>7,} {r[2]:>7,} {r[3]:>7,} {r[4]:>7,}")
    print(f"\nwrote {os.path.relpath(OUT, os.path.join(HERE, '..', '..'))}")


if __name__ == "__main__":
    main()
