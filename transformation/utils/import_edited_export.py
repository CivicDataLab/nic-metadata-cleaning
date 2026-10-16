#!/usr/bin/env python3
"""
Write hand edits made in an exported dublin_core_remaining workbook back to the DB.

export_dublin_core_remaining.py writes the table to xlsx; people then edit
merge_method / Collection / merge_add_columns etc. in a spreadsheet. This script
diffs that workbook against the live table on nid and writes back ONLY a
whitelist of columns (--columns). Everything else in the workbook is ignored,
because the export itself is lossy: cells over Excel's 32767-char limit were
truncated and illegal XML characters stripped, so a blind write-back would
corrupt long columns such as "Conforms To".

Cleaning applied to edited values:
  * leading/trailing whitespace stripped, empty cell -> NULL
  * --rename OLD=>NEW rewrites a Collection name on every row that ends up
    carrying OLD (edited or not), e.g. to settle a case-variant pair

Safety, in order:
  1. Default is a dry run: DB opened read_only=True, diff printed only.
  2. Workbook nids missing from the table, and table nids missing from the
     workbook, are reported. A column value in the DB longer than 32767 chars is
     never overwritten (the workbook copy of it is truncated).
  3. --apply takes a CSV snapshot of nid + the whitelisted columns to
     data/backups/ first.
  4. One UPDATE ... FROM a temp table (never per-row UPDATEs, which corrupted
     this DB twice).
  5. Post-write: row count must be unchanged and every changed row must read
     back with its new values; curate rows without a Collection and collections
     spanning more than one merge_method are reported.

Usage:
  python3 import_edited_export.py dublin_core_remaining_20260914_141424.xlsx
  python3 import_edited_export.py FILE.xlsx --rename "Town Amenities by district 2011=>Town Amenities by District 2011" --apply
"""

import argparse
import csv
import datetime as dt
import sys
from collections import Counter
from pathlib import Path

import duckdb

HERE = Path(__file__).resolve().parent
DB_PATH = str(HERE.parent / "metadata.db")
TABLE = "dublin_core_remaining"
DATA_DIR = HERE.parent.parent / "data"
DEFAULT_COLUMNS = ["merge_method", "Collection", "merge_add_columns", "state"]
EXCEL_CELL_LIMIT = 32767


def norm(v):
    """Comparable, cleaned form of a cell: stripped string or None."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def load_workbook_edits(path: Path, columns: list[str]) -> dict[str, dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [c for c in ["nid", *columns] if c not in header]
    if missing:
        sys.exit(f"workbook has no column(s) {missing}; header starts {header[:12]}")
    idx = {c: header.index(c) for c in ["nid", *columns]}
    out, dup = {}, 0
    for row in rows:
        nid = norm(row[idx["nid"]]) if idx["nid"] < len(row) else None
        if nid is None:
            continue
        if nid in out:
            dup += 1
        out[nid] = {c: norm(row[i]) if i < len(row) else None for c, i in idx.items() if c != "nid"}
    wb.close()
    if dup:
        print(f"  ! {dup} duplicate nid rows in workbook (last one wins)")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--db", default=DB_PATH)
    ap.add_argument("--table", default=TABLE)
    ap.add_argument("--columns", default=",".join(DEFAULT_COLUMNS),
                    help=f"comma list of columns to write back (default {','.join(DEFAULT_COLUMNS)})")
    ap.add_argument("--rename", action="append", default=[],
                    help="OLD=>NEW Collection rename applied after the edits (repeatable)")
    ap.add_argument("--changes-out", default=None, help="write the per-cell change list to this csv")
    ap.add_argument("--apply", action="store_true", help="write (default is dry-run)")
    args = ap.parse_args()

    columns = [c.strip() for c in args.columns.split(",") if c.strip()]
    renames = {}
    for r in args.rename:
        if "=>" not in r:
            sys.exit(f"--rename needs OLD=>NEW, got {r!r}")
        old, new = (x.strip() for x in r.split("=>", 1))
        renames[old] = new
    if renames and "Collection" not in columns:
        sys.exit("--rename needs Collection among --columns")

    print(f"reading {args.workbook} ...")
    edits = load_workbook_edits(Path(args.workbook), columns)
    print(f"workbook rows : {len(edits)}")

    con = duckdb.connect(args.db, read_only=not args.apply)
    try:
        qcols = ", ".join(f'"{c}"' for c in columns)
        db = {str(r[0]): dict(zip(columns, r[1:]))
              for r in con.execute(f"SELECT nid, {qcols} FROM {args.table}").fetchall()}
        print(f"table rows    : {len(db)}")
        only_wb = [n for n in edits if n not in db]
        only_db = [n for n in db if n not in edits]
        print(f"nids only in workbook: {len(only_wb)}   only in table: {len(only_db)}")

        changes = []           # (nid, column, old, new)
        skipped_long = Counter()
        final = {}
        for nid, cur in db.items():
            new_row = dict(cur)
            wb = edits.get(nid)
            if wb:
                for c in columns:
                    old = cur[c]
                    if old is not None and len(str(old)) > EXCEL_CELL_LIMIT:
                        skipped_long[c] += 1
                        continue
                    if norm(old) != wb[c]:
                        new_row[c] = wb[c]
                    elif isinstance(old, str) and old != wb[c]:
                        new_row[c] = wb[c]        # whitespace-only difference: clean it
            if "Collection" in columns and new_row["Collection"] in renames:
                new_row["Collection"] = renames[new_row["Collection"]]
            for c in columns:
                if new_row[c] != cur[c]:
                    changes.append((nid, c, cur[c], new_row[c]))
            if new_row != cur:
                final[nid] = new_row

        per_col = Counter(c for _, c, _, _ in changes)
        print(f"\nrows to change: {len(final)}")
        for c in columns:
            print(f"  {c:<20} {per_col[c]:>6}")
        if skipped_long:
            print(f"  (not touched, DB value longer than the xlsx cell limit: {dict(skipped_long)})")

        for c in columns:
            trans = Counter((o, n) for _, col, o, n in changes if col == c)
            if not trans:
                continue
            print(f"\n  {c} - top transitions:")
            for (o, n), k in trans.most_common(12):
                print(f"    [{k:5d}] {str(o)[:70]!r} -> {str(n)[:70]!r}")

        if args.changes_out:
            with open(args.changes_out, "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                w.writerow(["nid", "column", "old", "new"])
                w.writerows(changes)
            print(f"\nchange list: {args.changes_out}")

        if not final:
            print("\nNothing to do.")
            return
        if not args.apply:
            print("\n[dry-run] No changes written. Re-run with --apply.")
            return

        ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        snap = DATA_DIR / "backups" / f"{args.table}_import_edits_{ts}.csv"
        snap.parent.mkdir(parents=True, exist_ok=True)
        con.execute(f"COPY (SELECT nid, {qcols} FROM {args.table}) TO '{snap}' (HEADER, DELIMITER ',')")
        print(f"\nPre-update snapshot : {snap}")
        total_before = con.execute(f"SELECT count(*) FROM {args.table}").fetchone()[0]

        con.execute("DROP TABLE IF EXISTS _tmp_import_edits")
        con.execute("CREATE TEMP TABLE _tmp_import_edits(nid VARCHAR, "
                    + ", ".join(f'"{c}" VARCHAR' for c in columns) + ")")
        con.executemany(f"INSERT INTO _tmp_import_edits VALUES ({', '.join('?' * (len(columns) + 1))})",
                        [(nid, *[row[c] for c in columns]) for nid, row in final.items()])
        sets = ", ".join(f'"{c}" = t."{c}"' for c in columns)
        con.execute(f"UPDATE {args.table} AS d SET {sets} FROM _tmp_import_edits AS t "
                    f"WHERE CAST(d.nid AS VARCHAR) = t.nid")

        fails = []
        total_after = con.execute(f"SELECT count(*) FROM {args.table}").fetchone()[0]
        if total_after != total_before:
            fails.append(f"row count changed {total_before} -> {total_after}")
        bad = con.execute(
            f"SELECT count(*) FROM {args.table} AS d JOIN _tmp_import_edits AS t "
            f"ON CAST(d.nid AS VARCHAR) = t.nid WHERE "
            + " OR ".join(f'd."{c}" IS DISTINCT FROM t."{c}"' for c in columns)).fetchone()[0]
        if bad:
            fails.append(f"{bad} rows do not read back with their new values")
        con.execute("DROP TABLE IF EXISTS _tmp_import_edits")

        audit = DATA_DIR / "backups" / f"{args.table}_import_edits_{ts}_changes.csv"
        with open(audit, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["nid", "column", "old", "new"])
            w.writerows(changes)
        print(f"Updated {len(final)} rows. Change list: {audit}")

        if fails:
            print("\nINVARIANT FAILURES:")
            for f in fails:
                print(f"  ! {f}")
            print(f"Revert from: {snap}")
            sys.exit(2)
        print("All invariants hold.")

        if "merge_method" in columns and "Collection" in columns:
            nocoll = con.execute(f"""SELECT count(*) FROM {args.table}
                                     WHERE merge_method='curate' AND "Collection" IS NULL""").fetchone()[0]
            multi = con.execute(f"""SELECT "Collection", list(DISTINCT merge_method) FROM {args.table}
                                    WHERE "Collection" IS NOT NULL GROUP BY 1
                                    HAVING count(DISTINCT merge_method) > 1""").fetchall()
            print(f"\nnotes: curate rows without a Collection: {nocoll}; "
                  f"collections spanning >1 merge_method: {len(multi)}")
            for name, mms in multi[:10]:
                print(f"   {name[:90]}  {mms}")
    finally:
        con.close()


if __name__ == "__main__":
    main()
